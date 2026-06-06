from datetime import date
from decimal import Decimal
from uuid import UUID

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.deps import get_db
from app.main import create_app
from app.models import AccountingLine, BankTransaction, Enterprise, Invoice, MatchRecord, MonthlyWorkPackage, TaxFilingDraft
from app.services.statement_service import generate_monthly_statement
from app.services.tax_service import calculate_vat_draft, export_tax_filing_draft, generate_tax_filing_draft


ORG = UUID("00000000-0000-0000-0000-000000000001")


def as_decimal(value) -> Decimal:
    return Decimal(str(value))


def make_package(db_session):
    enterprise = Enterprise(
        organization_id=ORG,
        name="苏州报表测试企业",
        unified_social_credit_code="91320500STAT000001",
        taxpayer_type="GENERAL",
        industry="批发业",
    )
    db_session.add(enterprise)
    db_session.flush()
    package = MonthlyWorkPackage(
        organization_id=ORG,
        enterprise_id=enterprise.id,
        period_year=2026,
        period_month=5,
    )
    db_session.add(package)
    db_session.commit()
    return package


def add_invoice(
    db_session,
    package,
    *,
    direction: str,
    number: str,
    amount: Decimal,
    tax_amount: Decimal,
    confirmed: bool = True,
):
    invoice = Invoice(
        organization_id=ORG,
        monthly_work_package_id=package.id,
        invoice_direction=direction,
        invoice_number=number,
        invoice_date=date(2026, 5, 10),
        amount=amount,
        tax_amount=tax_amount,
        total_amount=amount + tax_amount,
    )
    db_session.add(invoice)
    db_session.flush()
    if confirmed:
        transaction = BankTransaction(
            organization_id=ORG,
            monthly_work_package_id=package.id,
            transaction_date=date(2026, 5, 12),
            summary=f"匹配{number}",
            debit_amount=invoice.total_amount if direction == "INPUT" else Decimal("0"),
            credit_amount=invoice.total_amount if direction == "OUTPUT" else Decimal("0"),
        )
        db_session.add(transaction)
        db_session.flush()
        db_session.add(
            MatchRecord(
                organization_id=ORG,
                monthly_work_package_id=package.id,
                bank_transaction_id=transaction.id,
                invoice_id=invoice.id,
                match_method="AUTO_EXACT",
                confidence=95,
                explanation="测试确认",
                confirmation_status="AUTO_CONFIRMED",
            )
        )
    return invoice


def test_calculate_vat_draft_from_invoice_totals():
    draft = calculate_vat_draft(
        output_amount=Decimal("100000"),
        output_tax=Decimal("13000"),
        input_amount=Decimal("40000"),
        input_tax=Decimal("5200"),
    )

    assert draft["vat_payable"] == Decimal("7800.00")
    assert draft["surcharge_estimate"] == Decimal("936.00")


def test_generate_monthly_statement_from_confirmed_data(db_session):
    package = make_package(db_session)
    add_invoice(db_session, package, direction="OUTPUT", number="OUT-1", amount=Decimal("100000"), tax_amount=Decimal("13000"))
    add_invoice(db_session, package, direction="INPUT", number="IN-1", amount=Decimal("40000"), tax_amount=Decimal("5200"))
    add_invoice(
        db_session,
        package,
        direction="OUTPUT",
        number="OUT-UNMATCHED",
        amount=Decimal("30000"),
        tax_amount=Decimal("3900"),
        confirmed=False,
    )
    db_session.add(
        AccountingLine(
            organization_id=ORG,
            monthly_work_package_id=package.id,
            source_type="BANK_TRANSACTION",
            source_id="expense-1",
            business_type="BANK_FEE",
            direction="EXPENSE",
            amount=Decimal("800"),
            tax_amount=Decimal("0"),
            include_category="EXPENSE",
            confirmation_status="CONFIRMED",
        )
    )
    db_session.add(
        BankTransaction(
            organization_id=ORG,
            monthly_work_package_id=package.id,
            transaction_date=date(2026, 5, 20),
            summary="额外付款",
            debit_amount=Decimal("2000"),
            credit_amount=Decimal("0"),
        )
    )
    db_session.commit()

    statement = generate_monthly_statement(db_session, monthly_work_package_id=package.id)

    assert as_decimal(statement.estimated_income_statement["revenue"]) == Decimal("100000.00")
    assert as_decimal(statement.estimated_income_statement["cost"]) == Decimal("40000.00")
    assert as_decimal(statement.estimated_income_statement["expense"]) == Decimal("800.00")
    assert as_decimal(statement.estimated_income_statement["operating_profit"]) == Decimal("59200.00")
    assert as_decimal(statement.estimated_balance_sheet["cash_net_movement"]) == Decimal("65800.00")


def test_latest_statement_html_endpoint_renders_online_preview(db_session):
    package = make_package(db_session)
    add_invoice(db_session, package, direction="OUTPUT", number="OUT-1", amount=Decimal("100000"), tax_amount=Decimal("13000"))
    db_session.commit()
    generate_monthly_statement(db_session, monthly_work_package_id=package.id)
    app = create_app(init_db_on_startup=False)

    def override_get_db():
        yield db_session

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app, raise_server_exceptions=False)

    response = client.get(f"/api/monthly-packages/{package.id}/statements/latest/html")

    assert response.status_code == 200
    assert "苏州报表测试企业" in response.text
    assert "每月账目与报表" in response.text
    assert "利润表估算" in response.text
    assert "营业收入" in response.text
    assert "revenue" not in response.text


def test_generate_tax_filing_draft_persists_invoice_totals(db_session):
    package = make_package(db_session)
    add_invoice(db_session, package, direction="OUTPUT", number="OUT-1", amount=Decimal("100000"), tax_amount=Decimal("13000"))
    add_invoice(db_session, package, direction="INPUT", number="IN-1", amount=Decimal("40000"), tax_amount=Decimal("5200"))
    db_session.commit()

    draft = generate_tax_filing_draft(db_session, monthly_work_package_id=package.id)

    assert as_decimal(draft.data["vat_payable"]) == Decimal("7800.00")
    assert as_decimal(draft.data["surcharge_estimate"]) == Decimal("936.00")
    assert draft.status == "DRAFT"
    assert db_session.query(TaxFilingDraft).count() == 1


def test_generate_tax_filing_draft_uses_imported_invoice_totals_before_matching(db_session):
    package = make_package(db_session)
    add_invoice(
        db_session,
        package,
        direction="OUTPUT",
        number="OUT-UNMATCHED",
        amount=Decimal("100000"),
        tax_amount=Decimal("13000"),
        confirmed=False,
    )
    add_invoice(
        db_session,
        package,
        direction="INPUT",
        number="IN-UNMATCHED",
        amount=Decimal("40000"),
        tax_amount=Decimal("5200"),
        confirmed=False,
    )
    db_session.commit()

    draft = generate_tax_filing_draft(db_session, monthly_work_package_id=package.id)

    assert as_decimal(draft.data["output_amount"]) == Decimal("100000.00")
    assert as_decimal(draft.data["output_tax"]) == Decimal("13000.00")
    assert as_decimal(draft.data["input_amount"]) == Decimal("40000.00")
    assert as_decimal(draft.data["input_tax"]) == Decimal("5200.00")
    assert as_decimal(draft.data["vat_payable"]) == Decimal("7800.00")
    assert draft.data["warnings"] == ["未匹配发票2张"]


def test_generate_tax_filing_draft_normalizes_imported_invoice_signs(db_session):
    package = make_package(db_session)
    add_invoice(
        db_session,
        package,
        direction="OUTPUT",
        number="OUT-NEGATIVE-SIGN",
        amount=Decimal("-100000"),
        tax_amount=Decimal("-13000"),
        confirmed=False,
    )
    add_invoice(
        db_session,
        package,
        direction="INPUT",
        number="IN-NEGATIVE-SIGN",
        amount=Decimal("-40000"),
        tax_amount=Decimal("-5200"),
        confirmed=False,
    )
    db_session.commit()

    draft = generate_tax_filing_draft(db_session, monthly_work_package_id=package.id)

    assert as_decimal(draft.data["output_amount"]) == Decimal("100000.00")
    assert as_decimal(draft.data["output_tax"]) == Decimal("13000.00")
    assert as_decimal(draft.data["input_amount"]) == Decimal("40000.00")
    assert as_decimal(draft.data["input_tax"]) == Decimal("5200.00")
    assert as_decimal(draft.data["vat_payable"]) == Decimal("7800.00")


def test_tax_filing_draft_html_endpoint_renders_online_preview(db_session):
    package = make_package(db_session)
    add_invoice(
        db_session,
        package,
        direction="OUTPUT",
        number="OUT-UNMATCHED",
        amount=Decimal("100000"),
        tax_amount=Decimal("13000"),
        confirmed=False,
    )
    db_session.commit()
    draft = generate_tax_filing_draft(db_session, monthly_work_package_id=package.id)
    app = create_app(init_db_on_startup=False)

    def override_get_db():
        yield db_session

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app, raise_server_exceptions=False)

    response = client.get(f"/api/tax-drafts/{draft.id}/html")

    assert response.status_code == 200
    assert "苏州报表测试企业" in response.text
    assert "增值税申报草稿" in response.text
    assert "销项销售额" in response.text
    assert "100000.00" in response.text


def test_update_tax_filing_draft_endpoint_recalculates_payable(db_session):
    package = make_package(db_session)
    add_invoice(
        db_session,
        package,
        direction="OUTPUT",
        number="OUT-UNMATCHED",
        amount=Decimal("100000"),
        tax_amount=Decimal("13000"),
        confirmed=False,
    )
    db_session.commit()
    draft = generate_tax_filing_draft(db_session, monthly_work_package_id=package.id)
    app = create_app(init_db_on_startup=False)

    def override_get_db():
        yield db_session

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app, raise_server_exceptions=False)

    response = client.patch(
        f"/api/tax-drafts/{draft.id}",
        json={
            "output_amount": "100000",
            "output_tax": "13000",
            "input_amount": "20000",
            "input_tax": "2600",
        },
    )

    assert response.status_code == 200
    assert response.json()["data"]["vat_payable"] == "10400.00"
    assert response.json()["data"]["surcharge_estimate"] == "1248.00"
    assert response.json()["data"]["warnings"] == ["未匹配发票1张"]


def test_export_tax_filing_draft_workbook_contains_copyable_rows(db_session, tmp_path):
    package = make_package(db_session)
    add_invoice(db_session, package, direction="OUTPUT", number="OUT-1", amount=Decimal("100000"), tax_amount=Decimal("13000"))
    add_invoice(db_session, package, direction="INPUT", number="IN-1", amount=Decimal("40000"), tax_amount=Decimal("5200"))
    add_invoice(
        db_session,
        package,
        direction="OUTPUT",
        number="OUT-UNMATCHED",
        amount=Decimal("30000"),
        tax_amount=Decimal("3900"),
        confirmed=False,
    )
    db_session.commit()
    draft = generate_tax_filing_draft(db_session, monthly_work_package_id=package.id)

    export_path = export_tax_filing_draft(db_session, draft_id=draft.id, output_dir=tmp_path)

    workbook = load_workbook(export_path)
    rows = list(workbook.active.iter_rows(values_only=True))
    assert rows[0] == ("字段", "金额", "说明")
    assert ("销项销售额", "130000.00", "本期销项不含税销售额") in rows
    assert ("本期应纳增值税", "11700.00", "销项税额减进项税额") in rows
    assert ("异常提醒", "未匹配发票1张", "导出前请人工核对") in rows
