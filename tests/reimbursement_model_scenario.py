"""Opt-in anonymous real-model check, isolated from business DB and original ledgers."""
from dataclasses import replace
from decimal import Decimal
from pathlib import Path
import argparse
import json
import sys

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app.autonomous_runtime import run
from app.bookkeeping_profile import snapshot
from app.config import Settings
from app.review_loop import REVIEW_POLICY_VERSION, effective_entries
from app.stage23_pilot import sha256
from app.travel_policy import POLICY_RULES


def main(output):
    output.mkdir(parents=True, exist_ok=False)
    book = openpyxl.Workbook(); bank=book.active; bank.title='银行'
    bank.append(['日期','收入','支出','收款人','摘要'])
    bank.append(['2026-08-05',0,266,'员工甲','差旅报销付款，按个人往来结算'])
    bank.append(['2026-08-06',0,300,'员工丙','差旅报销付款，票据暂未提供'])
    bills=book.create_sheet('费用票据'); bills.append(['票号','日期','金额','税额','价税合计','报销人','说明'])
    bills.append(['T1','2026-08-03',240,0,240,'员工甲','公司差旅车票；归属已经由报销单确认，本期首次入账'])
    bills.append(['T2','2026-08-03',80,0,80,'员工乙','公司差旅车票；归属已经由报销单确认，本期首次入账'])
    bills.append(['T3','2026-08-04',50,0,50,'','公司差旅车票，无法确定报销人，暂留待分配'])
    bills.append(['T0','2026-07-20',200,0,200,'员工甲','历史票据重复提供，7月已记管理费用/其他应付款甲200，本期不得重复记费'])
    opening=book.create_sheet('期初余额'); opening.append(['科目','名称','借方','贷方'])
    for row in [('1002','银行存款',1000,0),('224101','其他应付款—员工甲',0,200),
                ('224102','其他应付款—员工乙',100,0),('224103','其他应付款—员工丙',0,0),('4001','实收资本',0,900)]:
        opening.append(row)
    chart=book.create_sheet('科目表'); chart.append(['编码','名称'])
    for row in [('1002','银行存款'),('2241','其他应付款'),('224101','其他应付款—员工甲'),
                ('224102','其他应付款—员工乙'),('224103','其他应付款—员工丙'),
                ('560205','管理费用—差旅费'),('3103','本年利润'),('4001','实收资本')]:chart.append(row)
    facts=book.create_sheet('完整业务说明');facts.append(['说明'])
    facts.append(['本企业本期仅这些业务，无本期原账；全部是公司差旅报销，无借款、补贴、其他结算或重复付款。票据均为普通票，无抵扣进项。本期无销售/生产。'])
    facts.append(['所给期初是已记账正式余额。员工乙上月已付100但未取得票，本期80元票据归属员工乙已确认；员工丙是新报销往来，期初为零。三人之间无代收、转让或抵扣关系。'])
    facts.append(['按已确认企业口径：票据按个人挂其他应付款，实际报销付款冲个人往来；不要求单笔等额，余额下月继续。报销人不明只保留票据待分配，不猜人、不挡已知付款。全部资料已提供。'])
    source=output/'anonymous.xlsx';book.save(source)
    manifest={'period':'2026-08','sources':[{'id':'anonymous','filename':source.name,'path':'sources/anonymous.xlsx','sha256':sha256(source)}]}
    context={'review_loop_version':REVIEW_POLICY_VERSION,'statements':[],'materials_complete':True,
        'bookkeeping_profile':snapshot(),'travel_policy':{'id':'anonymous-travel','version':1,'rules':POLICY_RULES},
        'goal':'处理原件全部业务并按本期实际费用结转损益；先读完整历史说明及期初，生成待复核草稿。'}
    settings=replace(Settings.from_env(),autonomous_backend='claude_code',autonomous_execution_mode='docker',
        autonomous_container_image='finwise-backtest:20261001-ocr',autonomous_host_run_root=str(output/'autonomous-runs'),database_path=output/'test.db',autonomous_max_rounds=3)
    result=run(output/'autonomous-runs/run',manifest,{'anonymous':source},context,settings,lambda:None,
        lambda phase,round_no,message:print(phase,round_no,message,flush=True))
    (output/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    d=result['delivery'];accounts={a['person']:a for a in d['employee_reimbursements']}
    groups=effective_entries(d)
    checks={'mechanical':result['mechanical']['passed'],'independent_review':result['review']['verdict']=='pass',
        'balances':{p:a['closing_balance'] for p,a in accounts.items()}=={'员工甲':'174.00','员工乙':'-20.00','员工丙':'-300.00'},
        'invoice_first_expense_no_repeat':sum(Decimal(e['debit']) for rows in groups.values() for e in rows if e['code']=='560205')==320,
        'pending_owner_preserved':len(d['unassigned_reimbursement_invoices'])==1,
        'only_owner_question':len([f for f in d['findings'] if f['needs_user']])==1,
        'payments_reviewable':all(next(a for a in d['voucher_assessments'] if a['voucher']==p['voucher'])['status']=='READY_FOR_REVIEW'
            for a in accounts.values() for p in a['payments']),
        'no_formal_posting':result['formal_balance_impact'] is False}
    (output/'checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2))
    print(json.dumps(checks,ensure_ascii=False),flush=True)
    assert all(checks.values()),checks


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True)
    main(p.parse_args().output.resolve())
