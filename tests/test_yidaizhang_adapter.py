from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.integrations.yidaizhang import (
    TemplateProfile,
    YidaizhangBindingError,
    YidaizhangFormatError,
    bind_receipt_to_manifest,
    export_voucher_file,
    load_receipt,
)


def profile() -> TemplateProfile:
    return TemplateProfile.from_dict({
        "sheet": "凭证导入",
        "header_row": 1,
        "columns": {
            "voucher_number": "凭证号", "voucher_date": "日期", "summary": "摘要",
            "account_code": "科目编码", "account_name": "科目名称", "debit": "借方", "credit": "贷方",
        },
    })


def template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "凭证导入"
    sheet.append(["凭证号", "日期", "摘要", "科目编码", "科目名称", "借方", "贷方"])
    workbook.save(path)


def test_export_uses_explicit_profile_and_emits_binding_manifest(tmp_path):
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "out.xlsx"
    manifest_path = tmp_path / "manifest.json"
    template(template_path)
    export = {
        "object_id": "export-1",
        "data": {"package_id": "package-1", "voucher_version_id": "voucher-1", "voucher_version": 2,
                 "voucher_header": {"voucher_number": "记-001", "voucher_date": "2026-03-31", "summary": "采购业务"},
                 "lines": [
                     {"line_no": 1, "direction": "DEBIT", "account": "库存商品", "account_code": "1405", "amount": "100.00"},
                     {"line_no": 2, "direction": "CREDIT", "account": "银行存款", "account_code": "1002", "amount": "100.00"},
                 ]},
    }
    manifest = export_voucher_file(export=export, scope={"accounting_period_id": "2026-03"}, template_path=template_path, profile=profile(), output_path=output_path, manifest_path=manifest_path)
    assert manifest["export_id"] == "export-1"
    assert manifest["voucher_version"] == 2
    assert manifest["file"]["sha256"]
    workbook = load_workbook(output_path, data_only=True)
    values = list(workbook["凭证导入"].iter_rows(min_row=2, max_row=3, values_only=True))
    assert values[0][3:] == ("1405", "库存商品", 100, None)
    assert json.loads(manifest_path.read_text(encoding="utf-8"))["package_id"] == "package-1"


def test_receipt_requires_exact_binding_and_is_idempotency_ready(tmp_path):
    receipt_path = tmp_path / "receipt.json"
    receipt_path.write_text(json.dumps({
        "status": "IMPORTED", "external_batch": "YDZ-001", "export_id": "export-1",
        "voucher_version_id": "voucher-1", "voucher_version": 2, "period": "2026-03",
        "export_file_sha256": "file-hash",
    }), encoding="utf-8")
    payload = load_receipt(receipt_path)
    assert payload["payload_hash"]
    bound = bind_receipt_to_manifest(payload, {"export_id": "export-1", "voucher_version_id": "voucher-1", "voucher_version": 2, "period": "2026-03", "file": {"sha256": "file-hash"}})
    assert bound["adapter"] == "yidaizhang-file-v1"
    assert bound["manifest_file_sha256"]
    with pytest.raises(YidaizhangBindingError):
        bind_receipt_to_manifest(payload, {"export_id": "other", "voucher_version_id": "voucher-1", "voucher_version": 2})


def test_tabular_receipt_without_profile_is_blocked(tmp_path):
    path = tmp_path / "receipt.csv"
    path.write_text("status,external_batch\nIMPORTED,YDZ-001\n", encoding="utf-8")
    with pytest.raises(YidaizhangFormatError, match="明确列映射"):
        load_receipt(path)
