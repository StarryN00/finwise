from copy import deepcopy
from pathlib import Path

import openpyxl
import pytest

from app.autonomous_checks import settlement_match_errors, validate
from app.autonomous_runtime import prepare
from app.config import Settings
from app.review_loop import review_loop_approved, review_loop_errors
from app.stage23_pilot import read_tables, sha256


@pytest.fixture
def settlement(tmp_path):
    book = openpyxl.Workbook()
    book.active.title = 'Bank'
    book.active.append(['日期', '支出'])
    book.active.append(['2026-08-07', 1149.60])
    invoices = book.create_sheet('Invoices')
    invoices.append(['日期', '金额', '税额', '价税合计'])
    invoices.append(['2026-07-30', 920.47, 69.53, 990.00])
    book.create_sheet('Reimbursement').append(['已批准其他费用', 159.60])
    source = tmp_path / 'source.xlsx'
    book.save(source)
    manifest = {'period':'2026-08', 'sources':[{'id':'s', 'path':'source.xlsx', 'sha256':sha256(source)}]}
    bank = {'source_id':'s', 'sheet':'Bank', 'row':2}
    invoice = {'source_id':'s', 'sheet':'Invoices', 'row':2}
    delivery = {'period':'2026-08', 'mode':'SOURCE_BOOKKEEPING', 'source_roles':{'s':'付款、历史发票及报销组成'},
        'bank':[{'source':bank, 'date':'2026-08-07', 'income':'0.00', 'expense':'1149.60'}],
        'invoices':[{'source':invoice, 'date':'2026-07-30', 'kind':'进项', 'net':'920.47', 'tax':'69.53', 'gross':'990.00'}],
        'entries':[{'voucher':'001', 'date':'2026-08-07', 'code':code, 'name':name, 'debit':debit, 'credit':credit, 'source':bank}
                   for code,name,debit,credit in [('2241','员工往来','1149.60','0.00'),('1002','银行存款','0.00','1149.60')]],
        'voucher_assessments':[{'voucher':'001', 'status':'READY_FOR_REVIEW', 'reason':'先审阅有据付款，不确认费用', 'evidence':[bank]}],
        'findings':[{'id':'gap', 'title':'报销组成待确认', 'detail':'发票990.00，付款1149.60，差159.60；已确认付款，尚待对应报销归属',
                     'needs_user':True, 'blocks_confirmation':False, 'confirmation_reason':'当前仅为付款往来，未计费用或进项',
                     'affected_vouchers':['001'], 'evidence':[bank,invoice]}],
        'settlement_matches':[{'id':'m1', 'basis':'GROSS', 'status':'CANDIDATE', 'payment_source':bank,
                               'invoice_sources':[invoice], 'other_components':[], 'affected_vouchers':['001'],
                               'finding_ids':['gap'], 'reason':'历史发票与付款仅为候选，差额与业务归属待核实'}],
        'report':'原件支持付款，跨月不自动成为问题，尚不能确认费用或冲销不存在的前期应付。'*10}
    return tmp_path, manifest, delivery, read_tables(tmp_path, manifest)


def test_tax_equal_but_gross_different_cannot_be_complete(settlement):
    root, manifest, delivery, tables = settlement
    match = delivery['settlement_matches'][0]
    match['status'] = 'MATCHED'
    errors, checks = settlement_match_errors(delivery, tables)
    assert any('未覆盖付款' in e for e in errors)
    assert checks[0]['difference'] == '159.60'
    match['basis'] = 'TAX'
    assert any('税额不能' in e for e in settlement_match_errors(delivery, tables)[0])
    assert not validate(root, manifest, delivery, {})['passed']


def test_candidates_keep_gap_and_do_not_block_established_payment(settlement):
    root, manifest, delivery, _ = settlement
    result = validate(root, manifest, delivery, {})
    assert result['passed'], result['errors']
    assert result['settlement_checks'][0]['difference'] == '159.60'
    delivery['settlement_matches'][0]['finding_ids'] = []
    assert not validate(root, manifest, delivery, {})['passed']


def test_supported_reimbursement_components_can_cover_gross_gap(settlement):
    _, _, delivery, tables = settlement
    match = delivery['settlement_matches'][0]
    match['other_components'] = [{'amount':'159.60', 'source':{'source_id':'s', 'sheet':'Reimbursement', 'row':1},
                                 'reason':'报销原件列明其他费用'}]
    match['status'] = 'MATCHED'
    errors, checks = settlement_match_errors(delivery, tables)
    assert not errors and checks[0]['difference'] == '0.00'
    match['other_components'][0]['amount'] = '160.00'
    assert any('不能用付款差额造数' in e for e in settlement_match_errors(delivery, tables)[0])


def test_gross_is_computed_when_source_only_has_net_and_tax(settlement):
    _, _, delivery, tables = settlement
    tables['s']['Invoices'][0]['values'].pop()
    tables['s']['Invoices'][1]['values'].pop()
    errors, checks = settlement_match_errors(delivery, tables)
    assert not errors and checks[0]['invoice_gross'] == '990.00'


def test_july_invoice_can_match_august_payment_without_period_question(settlement):
    _, _, delivery, tables = settlement
    delivery['bank'][0]['expense'] = '990.00'
    tables['s']['Bank'][1]['values'][1] = 990.00
    match = delivery['settlement_matches'][0]
    match.update(status='MATCHED', finding_ids=[], reason='已核对业务身份和不重复，发票与付款跨月')
    assert not settlement_match_errors(delivery, tables)[0]
    duplicate = deepcopy(match); duplicate['id'] = 'm2'
    delivery['settlement_matches'].append(duplicate)
    assert any('不能重复完整匹配' in e for e in settlement_match_errors(delivery, tables)[0])


def test_cannot_forge_gross_or_reuse_payment_as_gap_evidence(settlement):
    _, _, delivery, tables = settlement
    invoice = delivery['invoices'][0]
    invoice.update(gross='1149.60', net='1080.07')
    assert any('价税合计须与原件一致' in e for e in settlement_match_errors(delivery, tables)[0])
    invoice.update(gross='990.00', net='920.47')
    delivery['settlement_matches'][0]['other_components'] = [{'source':delivery['bank'][0]['source'],
                                                            'amount':'159.60', 'reason':'用差额猜测无票费用'}]
    assert settlement_match_errors(delivery, tables)[0]


def test_independent_review_must_explicitly_cover_candidate_links(settlement):
    _, _, delivery, _ = settlement
    assert not review_loop_approved(delivery, {})
    assert review_loop_approved(delivery, {'reviewed_settlement_matches':['m1']})
    assert not review_loop_approved(delivery, {'reviewed_settlement_matches':['m1','invented']})
    delivery.pop('settlement_matches')
    assert any('关联清单' in e for e in review_loop_errors(delivery, {'review_loop_version':4}))
    assert not any('关联清单' in e for e in review_loop_errors(delivery, {'review_loop_version':3}))


def test_frozen_agent_checker_enforces_same_rules(settlement):
    root, manifest, delivery, _ = settlement
    settings = Settings(root=Path(__file__).resolve().parents[1], database_path=root/'db', storage_path=root/'data')
    workspace, _ = prepare(root/'run', manifest, {'s':root/'source.xlsx'}, {}, settings)
    namespace = {}
    exec((workspace/'tools/checker.py').read_text(), namespace)
    delivery['settlement_matches'][0]['status'] = 'MATCHED'
    result = namespace['validate'](workspace, manifest, delivery, {})
    assert not result['passed'] and any('未覆盖付款' in e for e in result['errors'])
