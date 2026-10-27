"""Opt-in real-model agency-bookkeeping examples; synthetic files, no business DB writes."""
import argparse
from dataclasses import replace
from decimal import Decimal
import json
from pathlib import Path
import sys

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app.autonomous_runtime import run
from app.config import Settings
from app.review_loop import REVIEW_POLICY_VERSION, effective_entries
from app.stage23_pilot import sha256


def verify_behavior(result):
    """Check the fixed synthetic facts, independently of the model's prose verdict."""
    delivery = result.get('delivery', {})
    review = result.get('review', {})
    assessments = delivery.get('voucher_assessments', [])
    findings = delivery.get('findings', [])
    entries = effective_entries(delivery)
    matches = delivery.get('settlement_matches', [])
    checks, errors = {}, {}

    def check(name, predicate):
        try:
            checks[name] = bool(predicate())
        except (ArithmeticError, KeyError, TypeError, ValueError) as exc:
            checks[name] = False
            errors[name] = str(exc)

    def money(value):
        amount = Decimal(str(value))
        if not amount.is_finite():
            raise ValueError('Amount must be finite')
        return amount

    def source(ref, sheet, row):
        return (isinstance(ref, dict) and ref.get('source_id') == 'synthetic'
                and ref.get('sheet') == sheet and ref.get('row') == row)

    def bank_source(ref, row):
        return source(ref, '银行回单', row)

    def assessment(name):
        rows = [a for a in assessments if a.get('voucher') == name]
        return rows[0] if len(rows) == 1 else {}

    def related_findings(name, row):
        ids = set(assessment(name).get('finding_ids', []))
        return [f for f in findings if name in f.get('affected_vouchers', []) or f.get('id') in ids
                or any(bank_source(ref, row) for ref in f.get('evidence', []))]

    def human_question(name, row):
        return any(f.get('needs_user') is True
                   and any(bank_source(ref, row) for ref in f.get('evidence', []))
                   and bool(f.get('request', {}).get('question') or f.get('decision', {}).get('question'))
                   for f in related_findings(name, row))

    def numbers_for_source(name, row):
        numbers = {name}
        numbers.update(a['voucher'] for a in assessments
                       if any(bank_source(ref, row) for ref in a.get('evidence', [])))
        numbers.update(number for number, rows in entries.items()
                       if any(bank_source(e.get('source'), row) for e in rows))
        for match in matches:
            if bank_source(match.get('payment_source'), row):
                numbers.update(match.get('affected_vouchers', []))
        return numbers

    def case_entries(name, row):
        return [entry for number in numbers_for_source(name, row) for entry in entries.get(number, [])]

    def ready_for_source(name, row):
        numbers = numbers_for_source(name, row)
        return any(a.get('status') == 'READY_FOR_REVIEW' and a.get('voucher') in numbers
                   for a in assessments) or bool(numbers & set(review.get('ready_vouchers', [])))

    def payment(row, amount):
        rows = [b for b in delivery.get('bank', []) if bank_source(b.get('source'), row)]
        return len(rows) == 1 and money(rows[0]['expense']) == money(amount)

    def total(name, side, code=None, row=None):
        rows = case_entries(name, row) if row is not None else entries.get(name, [])
        return sum((money(e[side]) for e in rows
                    if code is None or e.get('code') == code), Decimal('0'))

    check('independent_review_passed', lambda: review.get('verdict') == 'pass'
          and not review.get('defects') and bool(review.get('coverage')))
    check('mechanical_checks_passed', lambda: result.get('mechanical', {}).get('passed') is True)
    check('formal_balance_impact_is_false', lambda: result.get('formal_balance_impact') is False)
    check('current_review_contract', lambda: result.get('review_policy_version') == REVIEW_POLICY_VERSION)
    normal = [('正常宽带月结', 2, '371.30'), ('历史近似金额', 3, '120.20'),
              ('普通办公采购', 4, '200.00'), ('前期应付结算', 5, '300.00')]
    for name, row, amount in normal:
        check(name + ':ready_without_human_question', lambda name=name, row=row:
              assessment(name).get('status') == 'READY_FOR_REVIEW'
              and name in review.get('ready_vouchers', [])
              and any(bank_source(ref, row) for ref in assessment(name).get('evidence', []))
              and not any(f.get('needs_user') is True for f in related_findings(name, row)))
        check(name + ':source_and_amount', lambda name=name, row=row, amount=amount:
              payment(row, amount) and total(name, 'debit', row=row) == money(amount)
              and total(name, 'credit', row=row) == money(amount)
              and total(name, 'credit', '1002', row=row) == money(amount))

    check('prior_payable_settlement_only', lambda:
          total('前期应付结算', 'debit', '220201', row=5) == Decimal('300')
          and total('前期应付结算', 'credit', '1002', row=5) == Decimal('300')
          and all((e.get('code') == '220201' and money(e['credit']) == 0)
                  or (e.get('code') == '1002' and money(e['debit']) == 0)
                  for e in case_entries('前期应付结算', 5)))
    check('duplicate_payment_not_ready_as_expense', lambda:
          payment(6, '500') and human_question('真实重复', 6)
          and not (ready_for_source('真实重复', 6)
                   and any(e.get('code', '').startswith('6') and money(e['debit']) > 0
                           for e in case_entries('真实重复', 6))))
    check('missing_invoice_requires_row7_question', lambda:
          payment(7, '400') and human_question('缺票且用途不明', 7)
          and not ready_for_source('缺票且用途不明', 7))
    gap_matches = [m for m in matches if bank_source(m.get('payment_source'), 8)
                   and '票款差额' in m.get('affected_vouchers', [])]
    check('amount_gap_remains_candidate_with_question', lambda:
          payment(8, '630') and assessment('票款差额').get('status') == 'NEEDS_EVIDENCE'
          and not ready_for_source('票款差额', 8) and human_question('票款差额', 8)
          and bool(gap_matches) and all(m.get('status') == 'CANDIDATE' for m in gap_matches))
    check('amount_gap_invoice_stays_600', lambda:
          len([i for i in delivery.get('invoices', []) if source(i.get('source'), '发票账单', 7)
               and money(i['gross']) == Decimal('600')]) == 1
          and all(any(source(ref, '发票账单', 7) for ref in m.get('invoice_sources', [])) for m in gap_matches))
    check('amount_gap_does_not_invent_30_component', lambda:
          all(m.get('other_components') == [] for m in gap_matches)
          and sum((money(e['debit']) for e in case_entries('票款差额', 8)
                   if e.get('code') != '1002'), Decimal('0')) <= Decimal('600'))
    check('cross_year_requires_question', lambda:
          payment(9, '700') and human_question('跨年费用', 9)
          and not ready_for_source('跨年费用', 9))
    check('cross_year_has_no_unsupported_treatment_option', lambda:
          all(option.get('kind') != 'ACCOUNTING_CHOICE' and not option.get('recommended')
              for finding in related_findings('跨年费用', 9)
              for option in finding.get('decision', {}).get('options', [])))
    check('cross_year_has_no_expense_without_treatment_basis', lambda:
          not any(e.get('code', '').startswith('6') and money(e['debit']) > 0
                  for e in case_entries('跨年费用', 9)))
    old = [f for f in findings if f.get('id') == 'old-month']
    check('old_month_question_preserved_as_resolved', lambda:
          len(old) == 1 and old[0].get('needs_user') is False
          and any(bank_source(ref, 2) for ref in old[0].get('evidence', [])))
    failures = [name for name, passed in checks.items() if not passed]
    return {'passed': not failures, 'checks': checks, 'failures': failures, 'errors': errors}


def verify_existing(path):
    result_path = path / 'result.json' if path.is_dir() else path
    try:
        result = json.loads(result_path.read_text())
        report = verify_behavior(result)
    except (OSError, ValueError, TypeError, AttributeError) as exc:
        report = {'passed': False, 'checks': {'readable_result': False},
                  'failures': ['readable_result'], 'errors': {'readable_result': str(exc)}}
    report['result_path'] = str(result_path.resolve())
    report_path = result_path.parent / 'behavior-checks.json'
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps(report, ensure_ascii=False), flush=True)
    return 0 if report['passed'] else 1


def main(output):
    output.mkdir(parents=True, exist_ok=False)
    book = openpyxl.Workbook(); bank = book.active; bank.title = '银行回单'
    bank.append(['日期', '场景', '摘要', '收入', '支出', '收款人'])
    cases = [
        ['2026-08-07', '正常宽带月结', '7月手机宽带费；账单B-JUL', 0, 371.30, '甲通信公司'],
        ['2026-08-08', '历史近似金额', '7月水费；账单W-JUL', 0, 120.20, '乙水务公司'],
        ['2026-08-09', '普通办公采购', '办公用打印纸；发票P-AUG', 0, 200, '丙文具公司'],
        ['2026-08-10', '前期应付结算', '结算7月电费E-JUL', 0, 300, '丁电力公司'],
        ['2026-08-11', '真实重复', '再次支付同一张发票DUP-01的同一业务，无新增业务', 0, 500, '戊通信公司'],
        ['2026-08-12', '缺票且用途不明', '员工报销', 0, 400, '员工王甲'],
        ['2026-08-13', '票款差额', '7月宽带T-JUL，报销中630元组成不明', 0, 630, '员工李乙'],
        ['2026-08-14', '跨年费用', '2025年12月通信费Y-DEC，历史未记，处理方法未明确', 0, 700, '己通信公司'],
    ]
    for row in cases: bank.append(row)
    invoices = book.create_sheet('发票账单')
    invoices.append(['开票日期', '场景', '票号', '服务月份', '项目', '金额', '税额', '价税合计', '销售方'])
    for row in [
        ['2026-08-01', '正常宽带月结', 'B-JUL', '2026-07', '公司宽带及手机费', 371.30, 0, 371.30, '甲通信公司'],
        ['2026-08-01', '历史近似金额', 'W-JUL', '2026-07', '公司水费', 120.20, 0, 120.20, '乙水务公司'],
        ['2026-08-09', '普通办公采购', 'P-AUG', '2026-08', '公司办公打印纸', 200, 0, 200, '丙文具公司'],
        ['2026-07-31', '前期应付结算', 'E-JUL', '2026-07', '公司电费', 300, 0, 300, '丁电力公司'],
        ['2026-07-02', '真实重复', 'DUP-01', '2026-06', '公司通信费', 500, 0, 500, '戊通信公司'],
        ['2026-08-01', '票款差额', 'T-JUL', '2026-07', '公司宽带费', 600, 0, 600, '庚通信公司'],
        ['2025-12-31', '跨年费用', 'Y-DEC', '2025-12', '公司通信费', 700, 0, 700, '己通信公司'],
    ]: invoices.append(row)
    history = book.create_sheet('全部相关前期账')
    history.append(['日期', '凭证号', '摘要及票号', '借方科目', '贷方科目', '金额', '本笔期末未结余额'])
    for row in [
        ['2026-07-07', '7月记005', '6月通信费B-JUN；票据与付款已结清', '660202 办公费', '1002 银行存款', 371.10, 0],
        ['2026-07-08', '7月记006', '办公费；历史摘要未说明具体项目', '660202 办公费', '1002 银行存款', 120, 0],
        ['2026-07-31', '7月记020', '7月电费E-JUL已计提未付', '660202 办公费', '220201 应付账款—丁电力', 300, 300],
        ['2026-07-02', '7月记003', '通信费DUP-01已入费用且全额付款', '660202 办公费', '1002 银行存款', 500, 0],
    ]: history.append(row)
    history.append(['说明', '本表为全部相关历史', '未另确认常规跨月补记政策；除E-JUL外无本次付款对应未结应付，Y-DEC前期未记；其余本月票号未见已报销记录。'])
    accounts = book.create_sheet('科目表'); accounts.append(['编码', '名称'])
    for row in [('1002', '银行存款'), ('660202', '管理费用—办公费'), ('220201', '应付账款—丁电力'), ('122101', '其他应收款—戊通信公司')]:
        accounts.append(row)
    source = output/'synthetic.xlsx'; book.save(source)
    manifest = dict(period='2026-08', sources=[dict(id='synthetic', filename=source.name,
                    path='sources/synthetic.xlsx', sha256=sha256(source))])
    context = dict(review_loop_version=REVIEW_POLICY_VERSION, statements=[], materials_complete=True,
                   goal='按代理记账日常口径独立查证全部8个场景。凭证号使用场景名称便于核对，不能把场景名称当作处理结论。没有本期原账。',
                   prior_findings=[dict(id='old-month', title='7月宽带费未见计提，是否在8月补记', needs_user=True,
                       detail='旧结果因未另提供跨月政策要求人工选择；请按当前契约重新从原件判断。')])
    settings = replace(Settings.from_env(), autonomous_backend='claude_code', autonomous_max_rounds=3)
    progress = lambda phase, round_no, message: print(json.dumps(dict(phase=phase, round=round_no, message=message), ensure_ascii=False), flush=True)
    result = run(output/'run', manifest, {'synthetic':source}, context, settings, lambda *args: None, progress)
    (output/'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
    print('MODEL_RESULT_SAVED', flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument('--output', type=Path)
    mode.add_argument('--verify-existing', type=Path, help='Verify an existing result.json or its directory without calling the model')
    args = parser.parse_args()
    if args.verify_existing is not None:
        raise SystemExit(verify_existing(args.verify_existing))
    main(args.output)
