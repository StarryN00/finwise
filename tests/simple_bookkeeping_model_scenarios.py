"""Opt-in real-model checks for receipts and unnecessary supplier follow-up."""
import argparse
from dataclasses import replace
from decimal import Decimal
import json
from pathlib import Path
import sys

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app.autonomous_runtime import run
from app.config import Settings
from app.review_loop import REVIEW_POLICY_VERSION, effective_entries
from app.stage23_pilot import sha256


def verify(result):
    delivery = result['delivery']
    entries = effective_entries(delivery)
    assessments = {a['voucher']: a for a in delivery['voucher_assessments']}
    findings = delivery['findings']
    checks = {
        'review_pass': result['review']['verdict'] == 'pass' and not result['review'].get('defects'),
        'mechanical_pass': result['mechanical']['passed'],
        'contract_version': result['review_policy_version'] == REVIEW_POLICY_VERSION,
        'draft_only': result['formal_balance_impact'] is False,
    }
    def total(voucher, code, side):
        return sum((Decimal(str(e[side])) for e in entries.get(voucher, []) if e['code'] == code), Decimal(0))
    for voucher, debit_code, credit_code, amount in [
        ('S2', '1002', '112201', '2260'), ('S3', '1002', '2203', '300'),
        ('S4', '1002', '2241', '400'), ('S5', '1123', '1002', '37800')]:
        checks[voucher + '_correct_settlement'] = (
            total(voucher, debit_code, 'debit') == Decimal(amount)
            and total(voucher, credit_code, 'credit') == Decimal(amount)
            and sum(Decimal(str(e['credit'])) for e in entries.get(voucher, [])) == Decimal(amount))
        checks[voucher + '_no_extra_revenue'] = total(voucher, '6001', 'credit') == 0
    checks['S1_sales_and_tax'] = (total('S1', '1002', 'debit') == 1130
                                  and total('S1', '6001', 'credit') == 1000
                                  and total('S1', '222101', 'credit') == 130)
    for i in range(1, 6):
        voucher = f'S{i}'
        related = [f for f in findings if voucher in f.get('affected_vouchers', [])]
        checks[voucher + '_ready_without_question'] = (
            assessments.get(voucher, {}).get('status') == 'READY_FOR_REVIEW'
            and voucher in result['review'].get('ready_vouchers', [])
            and not any(f.get('needs_user') for f in related))
    unknown = [f for f in findings if f.get('needs_user') and (
        'S6' in f.get('affected_vouchers', []) or any(
            e.get('sheet') == '银行回单' and e.get('row') == 7 for e in f.get('evidence', [])))]
    checks['unknown_receipt_still_asks_specific_fact'] = bool(unknown) and all(
        f.get('request', {}).get('question') and f.get('blocks_confirmation') for f in unknown)
    checks['unknown_not_guessed_as_sales'] = total('S6', '6001', 'credit') == 0
    # Finding IDs may change between runs; match the retained note to the original payment.
    old = [f for f in findings if 'S5' in f.get('affected_vouchers', []) and any(
        e.get('source_id') == 'synthetic' and e.get('sheet') == '银行回单' and e.get('row') == 6
        for e in f.get('evidence', []))]
    checks['old_followup_retained_as_note'] = bool(old) and all(
        f.get('needs_user') is False and not f.get('decision') for f in old)
    checks['no_extra_sales_elsewhere'] = sum(
        Decimal(str(e['credit'])) for rows in entries.values() for e in rows if e['code'] == '6001') == 1000
    return {'passed': all(checks.values()), 'checks': checks}


def main(output):
    output.mkdir(parents=True, exist_ok=False)
    book = openpyxl.Workbook()
    bank = book.active; bank.title = '银行回单'
    bank.append(['编号', '日期', '摘要', '收入', '支出', '对方'])
    for row in [
        ['S1', '2026-08-08', '本月销售货款；销项票A01', 1130, 0, '客户甲'],
        ['S2', '2026-08-09', '付7月已开票货款A02', 2260, 0, '客户乙'],
        ['S3', '2026-08-10', '9月订单预付款，尚未发货，无销项发票', 300, 0, '客户丙'],
        ['S4', '2026-08-11', '客户临时借款给本公司，下月偿还，与销售无关', 400, 0, '客户丁'],
        ['S5', '2026-08-12', '采购设备预付款；目前仅付款，供应商未开票，无交付信息', 0, 37800, '供应商戊'],
        ['S6', '2026-08-13', '转账；无其他业务资料，未找到历史对应', 500, 0, '客户己'],
    ]: bank.append(row)
    sales = book.create_sheet('销售资料')
    sales.append(['编号', '发票', '月份', '金额', '销项税', '价税合计', '说明'])
    sales.append(['S1', 'A01', '2026-08', 1000, 130, 1130, '本月正常产品销售，已开票，本期未记收入，无退货或预收标记'])
    sales.append(['S2', 'A02', '2026-07', 2000, 260, 2260, '前期已记销售收入2000及销项税260，期末应收客户乙2260，本期收款结清'])
    notes = book.create_sheet('完整相关历史及口径')
    notes.append(['说明'])
    notes.append(['本企业按一般纳税人13%销售税率；银行S1至S6均为真实本期流水。S2除应收2260外无其他未结往来，其余本表本期业务前期均未记账。'])
    notes.append(['无本期原账。S1普通销售资料无冲突，不另提供客户交货验收过程。S5只记录有据预付款，不判断是否已经交付，也不承担经营跟进。S6真实性质尚不清楚。'])
    accounts = book.create_sheet('科目表'); accounts.append(['编码', '名称'])
    for row in [('1002', '银行存款'), ('6001', '主营业务收入'), ('222101', '应交税费—销项税额'),
                ('112201', '应收账款—客户乙'), ('2203', '预收账款'), ('2241', '其他应付款'), ('1123', '预付账款')]:
        accounts.append(row)
    source = output / 'synthetic.xlsx'; book.save(source)
    manifest = {'period': '2026-08', 'sources': [{'id': 'synthetic', 'filename': source.name,
                'path': 'sources/synthetic.xlsx', 'sha256': sha256(source)}]}
    context = {'review_loop_version': REVIEW_POLICY_VERSION, 'statements': [], 'materials_complete': True,
               'goal': '处理全部6笔业务；凭证号用银行编号S1至S6，编号只作定位。缺少必要事实不能编造分录。',
               'prior_findings': [{'id': 'old-followup', 'title': '供应商付款后需要确认是否交付',
                                   'detail': '旧版对S5仅作未来跟进，未影响当前付款记录。', 'needs_user': True}]}
    settings = replace(Settings.from_env(), autonomous_backend='claude_code', autonomous_max_rounds=3)
    result = run(output/'run', manifest, {'synthetic': source}, context, settings, lambda: None,
                 lambda phase, round_no, message: print(phase, round_no, message, flush=True))
    (output/'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
    report = verify(result)
    (output/'behavior-checks.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps(report, ensure_ascii=False), flush=True)
    return 0 if report['passed'] else 1


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument('--output', type=Path)
    mode.add_argument('--verify-existing', type=Path)
    args = parser.parse_args()
    if args.verify_existing:
        result_path = args.verify_existing / 'result.json'
        report = verify(json.loads(result_path.read_text()))
        (result_path.parent/'behavior-checks.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
        print(json.dumps(report, ensure_ascii=False))
        raise SystemExit(0 if report['passed'] else 1)
    raise SystemExit(main(args.output))
