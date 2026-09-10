"""Synthetic pure-function checks; no app, database, service or model fixtures."""
from copy import deepcopy
from decimal import Decimal, localcontext
import json

import pytest

from app.reconcile import VERSION, reconcile


def record(row, income="0", expense="0", balance=None, sheet="流水", account="a", **extra):
    return {"record_type": "BANK_TRANSACTION", "source_anchor": {"region": f"{sheet}!A{row}:H{row}"},
            "original_value": {"sheet": sheet, "row": row},
            "normalized_value": {"income": income, "expense": expense, "balance": balance,
                                 "bank_account_ref": account}, "period_check": "PASS", **extra}


def extraction(records, totals=None, sheet="流水"):
    refs = [] if totals is None else [
        {"row": 90, "reason": "BANK_TOTAL_HEADER", "values": ["总收入笔数", "总收入金额", "总支出笔数", "总支出金额"]},
        {"row": 91, "reason": "BANK_TOTAL_VALUES", "values": totals}]
    return {"records": records, "sheets": [{"sheet": sheet, "reference_rows": refs}]}


def check(extracted, code):
    return next(i for i in reconcile(extracted, "bank_statement")["items"] if i["code"] == code)


def test_four_columns_full_data_before_period_and_no_mutation():
    rows = [record(2, "10.10", record_type="RECEIPT"), record(3, expense="3.05", record_type="PAYMENT"),
            record(4, "0.20", period_check="PERIOD_EXCEPTION")]
    data = extraction(rows, [2, "10.30", 1, "3.05"])
    original = deepcopy(data)
    result = reconcile(data, "bank_statement")
    totals = [i for i in result["items"] if i["code"].startswith("BANK_TOTAL_")]
    assert len(totals) == 4
    assert all(i["status"] == "MATCH" for i in totals)
    assert result["overall"] == "MATCH"
    assert result["version"] == VERSION
    assert data == original
    assert "流水!A91:D91" in totals[0]["source_regions"]
    assert "流水!A4:H4" in totals[0]["source_regions"]
    json.dumps(result, allow_nan=False)


@pytest.mark.parametrize("column,code", [(0, "INCOME_COUNT"), (1, "INCOME_AMOUNT"),
                                         (2, "EXPENSE_COUNT"), (3, "EXPENSE_AMOUNT")])
def test_each_declared_measure_mismatch(column, code):
    totals = [1, "7", 1, "3"]
    totals[column] = "9"
    data = extraction([record(2, "7"), record(3, expense="3")], totals)
    assert check(data, "BANK_TOTAL_" + code)["status"] == "MISMATCH"
    assert reconcile(data, "bank_statement")["overall"] == "REVIEW"


@pytest.mark.parametrize("reverse", [False, True])
def test_balance_both_orders(reverse):
    rows = [record(2, "13", balance="113"), record(3, expense="4", balance="109"),
            record(4, "6", balance="115")]
    data = extraction(rows[::-1] if reverse else rows)
    item = check(data, "BANK_BALANCE_CHAIN")
    assert item["status"] == "MATCH"
    assert item["computed"]["valid_orders"] == ["REVERSE" if reverse else "FORWARD"]
    assert "开头是否漏读无法验证" in item["message"]
    assert reconcile(data, "bank_statement")["overall"] == "MATCH"


def test_balance_reports_both_failed_directions_with_evidence():
    data = extraction([record(2, "5", balance="105"), record(3, expense="3", balance="99")])
    item = check(data, "BANK_BALANCE_CHAIN")
    assert item["status"] == "MISMATCH"
    assert item["computed"]["FORWARD"]["balance"] == "102"
    assert item["computed"]["REVERSE"]["balance"] == "104"


@pytest.mark.parametrize("rows", [[], [record(2, "3", balance="100")],
    [record(2, "3"), record(3, "4")],
    [record(2, "3", balance="100"), record(3, "4"), record(4, "2", balance="106")]])
def test_missing_balance_or_no_independent_anchor_is_unattested(rows):
    result = reconcile(extraction(rows), "bank_statement")
    assert result["overall"] == "UNATTESTED"
    assert all(i["status"] == "NOT_AVAILABLE" for i in result["items"])


def test_prefix_is_suspect_only_even_when_totals_match():
    data = extraction([record(2, "2"), record(3, "3"), record(4, "5")], [3, "10", 0, "0"])
    original = deepcopy(data)
    item = check(data, "ARITHMETIC_TOTAL_PROBE")
    assert item["status"] == "SUSPECT"
    assert item["computed"]["prefix_sum"] == "5"
    assert reconcile(data, "bank_statement")["overall"] == "REVIEW"
    assert data == original


@pytest.mark.parametrize("rows", [[record(2, "2"), record(3, "2")],
    [record(2), record(3), record(4)], [record(2, "3"), record(3), record(4, "3")],
    [record(2, "2"), record(3, income=None), record(4, "3"), record(5, "5")]])
def test_prefix_does_not_invent_zero_or_suspect_single_contributor(rows):
    assert not any(i["code"] == "ARITHMETIC_TOTAL_PROBE" for i in reconcile(extraction(rows), "bank_statement")["items"])


@pytest.mark.parametrize("value", [None, "", "bad", True, "NaN", "Infinity", "1,23"])
def test_unknown_amount_is_not_zero(value):
    data = extraction([record(2, value)], [0, "0", 0, "0"])
    assert check(data, "BANK_TOTAL_INCOME_AMOUNT")["status"] == "NOT_AVAILABLE"
    assert check(data, "BANK_TOTAL_INCOME_COUNT")["status"] == "NOT_AVAILABLE"
    assert reconcile(data, "bank_statement")["overall"] == "UNATTESTED"


def test_exact_fractional_difference_not_hidden_by_display_format_or_decimal_context():
    data = extraction([record(2, "10000000000000000000000000.001")], [1, "10000000000000000000000000.00", 0, 0])
    data["records"][0]["original_value"]["number_format"] = "0.00"
    with localcontext() as context:
        context.prec = 6
        item = check(data, "BANK_TOTAL_INCOME_AMOUNT")
        assert context.prec == 6
    assert item["status"] == "MISMATCH"
    assert Decimal(item["computed"]) - Decimal(item["declared"]) == Decimal(".001")


def test_decimal_float_conversion_and_comma_declarations():
    data = extraction([record(2, 1000.1), record(3, Decimal("0.2"))], [2, "1,000.30", 0, 0])
    assert check(data, "BANK_TOTAL_INCOME_AMOUNT")["status"] == "MATCH"


def test_sheet_totals_and_account_chains_do_not_mix():
    data = extraction([record(2, "5", balance="105"), record(3, expense="2", balance="103"),
                       record(4, "8", balance="28", account="b"), record(5, expense="3", balance="25", account="b")],
                      [2, 13, 2, 5])
    other = extraction([record(2, "7", sheet="其他")], [1, 7, 0, 0], sheet="其他")
    data["records"].extend(other["records"])
    data["sheets"].extend(other["sheets"])
    result = reconcile(data, "bank_statement")
    assert result["overall"] == "MATCH"
    assert sum(i["status"] == "MATCH" and i["code"] == "BANK_BALANCE_CHAIN" for i in result["items"]) == 2


def test_partial_sheet_attestation_does_not_claim_whole_file_match():
    data = extraction([record(2, "3"), record(2, "4", sheet="其他")], [1, 3, 0, 0])
    assert reconcile(data, "bank_statement")["overall"] == "UNATTESTED"


def test_multiple_declarations_abstain_from_guessing_scope():
    data = extraction([record(2, "5")], [1, 5, 0, 0])
    data["sheets"][0]["reference_rows"].append({"reason": "BANK_TOTAL_VALUES", "row": 80, "values": [1, 5, 0, 0]})
    assert check(data, "BANK_TOTAL_SCOPE")["status"] == "NOT_AVAILABLE"
    assert reconcile(data, "bank_statement")["overall"] == "UNATTESTED"


def test_missing_sources_and_mixed_unknown_accounts_do_not_attest():
    data = extraction([record(2, "5", balance="105"), record(3, expense="2", balance="103", account=None)])
    assert reconcile(data, "bank_statement")["overall"] == "UNATTESTED"
    data = extraction([record(2, "5")], [1, 5, 0, 0])
    data["records"][0]["source_anchor"] = {}
    assert reconcile(data, "bank_statement")["overall"] == "UNATTESTED"


def test_nonbank_and_empty_old_boolean_checks_cannot_attest():
    data = {"records": [], "checks": {"statement_totals": True, "balance_continuity": True}}
    for kind in ("bank_statement", "purchase_invoices"):
        assert reconcile(data, kind)["overall"] == "UNATTESTED"


def test_negative_columns_and_nonintegral_counts_abstain():
    data = extraction([record(2, "-1")], [1, 1, 0, 0])
    assert check(data, "BANK_TOTAL_INCOME_AMOUNT")["status"] == "NOT_AVAILABLE"
    data = extraction([record(2, "1")], ["1.5", 1, 0, 0])
    assert check(data, "BANK_TOTAL_INCOME_COUNT")["status"] == "NOT_AVAILABLE"


def test_expense_prefix_in_middle_keeps_later_rows():
    data = extraction([record(2, expense="2"), record(3, expense="7"),
                       record(4, expense="9"), record(5, expense="1")])
    item = check(data, "ARITHMETIC_TOTAL_PROBE")
    assert item["computed"]["field"] == "expense"
    assert item["source_regions"][-1] == "流水!A4:H4"
    assert len(data["records"]) == 4


def test_prefix_never_sums_unrelated_accounts():
    data = extraction([record(2, "2"), record(3, "3"), record(4, "5", account="b")])
    assert not any(i["status"] == "SUSPECT" for i in reconcile(data, "bank_statement")["items"])


def test_dual_valid_balance_order_does_not_invent_direction():
    data = extraction([record(2, "5", balance="105"), record(3, expense="5", balance="100")])
    item = check(data, "BANK_BALANCE_CHAIN")
    assert item["status"] == "MATCH"
    assert item["computed"]["valid_orders"] == ["FORWARD", "REVERSE"]


def test_declared_sheet_without_details_prevents_whole_file_attestation():
    data = extraction([record(2, "5")], [1, 5, 0, 0])
    data["sheets"].extend(extraction([], [1, 7, 0, 0], sheet="遗漏页")["sheets"])
    assert reconcile(data, "bank_statement")["overall"] == "UNATTESTED"


def test_nonbank_row_is_not_silently_covered_by_bank_totals():
    data = extraction([record(2, "5"), record(3, "0", record_type="INVOICE")], [1, 5, 0, 0])
    assert reconcile(data, "bank_statement")["overall"] == "UNATTESTED"


def test_missing_sheet_is_not_assigned_to_only_summary():
    data = extraction([record(2, "5")], [1, 5, 0, 0])
    data["records"][0]["original_value"] = {}
    data["records"][0]["source_anchor"] = {}
    assert reconcile(data, "bank_statement")["overall"] == "UNATTESTED"


def test_twenty_thousand_rows_checked_without_truncating():
    data = extraction([record(i + 2, "1", balance=str(101 + i)) for i in range(20_000)],
                      [20_000, 20_000, 0, 0])
    result = reconcile(data, "bank_statement")
    assert result["overall"] == "MATCH"
    chain = next(i for i in result["items"] if i["code"] == "BANK_BALANCE_CHAIN")
    assert chain["computed"]["transitions"] == 19_999
    assert "流水!A20001:H20001" in chain["source_regions"]


def bank_workbook(rows):
    from io import BytesIO
    import openpyxl

    book = openpyxl.Workbook()
    book.active.title = "流水"
    book.active.append(["交易时间", "收入金额", "支出金额", "账户余额", "摘要"])
    for row in rows:
        book.active.append(row)
    content = BytesIO()
    book.save(content)
    book.close()
    return content.getvalue()


def test_tabular_checks_are_additive_and_use_records_after_footer_and_outside_period(monkeypatch):
    from app import reconcile as module
    from app.tabular import extract_workbook, ParseOptions

    content = bank_workbook([
        ["2026-01-01", 10, 0, 110],
        ["总收入笔数", "总收入金额", "总支出笔数", "总支出金额"],
        [1, 10, 1, 3], ["2026-02-01", 0, 3, 107]])
    options = ParseOptions(document_kind="bank_statement", bank_account_ref="a")
    result = extract_workbook(content, options, "2026-01")
    assert result["checks"]["overall"] == "MATCH"
    assert [r["period_check"] for r in result["records"]] == ["PASS", "PERIOD_EXCEPTION"]
    assert [r["source_anchor"]["row"] for r in result["records"]] == [2, 5]
    # Disable only the new checker to compare every pre-existing output field.
    monkeypatch.setattr(module, "reconcile", lambda *args: {})
    baseline = extract_workbook(content, options, "2026-01")
    assert {k: v for k, v in result.items() if k != "checks"} == {k: v for k, v in baseline.items() if k != "checks"}


def test_tabular_difference_is_parse_check_not_customer_evidence_or_fact_rewrite():
    from app.tabular import extract_workbook, ParseOptions

    result = extract_workbook(bank_workbook([
        ["2026-01-01", 10, 0, 110],
        ["总收入笔数", "总收入金额", "总支出笔数", "总支出金额"], [1, "10.01", 0, 0]]),
        ParseOptions(document_kind="bank_statement", bank_account_ref="a"), "2026-01")
    assert result["checks"]["overall"] == "REVIEW"
    assert result["errors"] == []
    fact = result["records"][0]
    assert fact["normalized_value"]["income"] == "10.00"
    assert fact["extraction_issues"] == []
    item = next(i for i in result["checks"]["items"] if i["status"] == "MISMATCH")
    assert "核查提取及原件" in item["message"]


def test_tabular_without_totals_or_balance_does_not_add_failure():
    from app.tabular import extract_workbook, ParseOptions

    result = extract_workbook(bank_workbook([["2026-01-01", 10, 0, None]]),
                              ParseOptions(document_kind="bank_statement", bank_account_ref="a"), "2026-01")
    assert result["checks"]["overall"] == "UNATTESTED"
    assert all(i["status"] == "NOT_AVAILABLE" for i in result["checks"]["items"])
    assert result["errors"] == [] and result["records"][0]["extraction_issues"] == []


def test_tabular_suspected_total_remains_a_record():
    from app.tabular import extract_workbook, ParseOptions

    result = extract_workbook(bank_workbook([["2026-01-01", 2, 0], ["2026-01-02", 3, 0],
                                           ["2026-01-03", 5, 0]]),
                             ParseOptions(document_kind="bank_statement", bank_account_ref="a"), "2026-01")
    assert result["checks"]["overall"] == "REVIEW"
    assert len(result["records"]) == 3
    assert [r["normalized_value"]["income"] for r in result["records"]] == ["2.00", "3.00", "5.00"]
    assert all(r["extraction_issues"] == [] for r in result["records"])


@pytest.mark.parametrize("state", ["hidden", "veryHidden"])
def test_xlsx_hidden_metadata_is_one_based_and_does_not_change_records(state):
    from io import BytesIO
    import openpyxl
    from app.tabular import read_workbook, extract_workbook, ParseOptions

    book = openpyxl.load_workbook(BytesIO(bank_workbook([["2026-01-01", 10, 0, 110]])))
    book.create_sheet("封面")["A1"] = "合成资料"
    before = BytesIO()
    book.save(before)
    page = book["流水"]
    page.sheet_state = state
    page.row_dimensions[2].hidden = True
    page.row_dimensions[20].hidden = True  # Metadata outside the data range matters too.
    page.column_dimensions.group("B", "D", hidden=True)
    page.column_dimensions.group("K", "L", hidden=True)
    page["B2"].number_format = "0.0000"
    after = BytesIO()
    book.save(after)
    book.close()
    sheet = read_workbook(after.getvalue())[0]
    assert sheet["sheet_state"] == state
    assert sheet["hidden_rows"] == [2, 20]
    assert sheet["hidden_columns"] == [2, 3, 4, 11, 12]
    assert sheet["rows"][1]["number_formats"][1] == "0.0000"
    options = ParseOptions(document_kind="bank_statement", bank_account_ref="a")
    assert extract_workbook(before.getvalue(), options, "2026-01")["records"] == extract_workbook(after.getvalue(), options, "2026-01")["records"]


def test_xlsx_visible_defaults_and_format_lists_align_with_values():
    from app.tabular import read_workbook

    sheet = read_workbook(bank_workbook([["2026-01-01", 10, None, 110]]))[0]
    assert sheet["sheet_state"] == "visible"
    assert sheet["hidden_rows"] == sheet["hidden_columns"] == []
    assert all(len(r["number_formats"]) == len(r["values"]) for r in sheet["rows"])


@pytest.mark.parametrize("format_table", [False, True])
def test_xls_layout_metadata_preserves_legacy_width(monkeypatch, format_table):
    from types import SimpleNamespace as NS
    import xlrd
    from app.tabular import read_workbook

    released = []
    cells = [NS(value="2026-01-01", ctype=xlrd.XL_CELL_TEXT), NS(value=10, ctype=xlrd.XL_CELL_NUMBER)]
    source = NS(name="流水", nrows=1, ncols=2, row=lambda r: cells)
    layout = NS(merged_cells=[], visibility=2, ncols=5,
                rowinfo_map={0: NS(hidden=1), 19: NS(hidden=1), 20: NS(hidden=0)},
                colinfo_map={1: NS(hidden=1), 4: NS(hidden=1)}, cell_xf_index=lambda r, c: c)

    def open_book(*, file_contents, on_demand, formatting_info=False):
        book = NS(nsheets=1, datemode=0, sheets=lambda: [source], sheet_by_index=lambda _: layout,
                  release_resources=lambda: released.append(formatting_info))
        if formatting_info and format_table:
            book.xf_list = [NS(format_key=0), NS(format_key=1)]
            book.format_map = {0: NS(format_str="General"), 1: NS(format_str="0.000")}
        return book

    monkeypatch.setattr(xlrd, "open_workbook", open_book)
    sheet = read_workbook(bytes.fromhex("d0cf11e0a1b11ae1") + b"synthetic")[0]
    assert sheet["rows"][0]["values"] == ["2026-01-01", 10]
    assert sheet["rows"][0]["number_formats"] == (["General", "0.000"] if format_table else [None, None])
    assert sheet["hidden_rows"] == [1, 20]
    assert sheet["hidden_columns"] == [2, 5]
    assert sheet["sheet_state"] == "veryHidden"
    assert released == [False, True]


def invoice_extraction(kind="purchase_invoices", reason="SUMMARY_REFERENCE", label="合计"):
    rows = []
    for row, net, tax, total in ((2, "0.10", "0.01", "0.11"), (3, "0.20", "0.02", "0.22")):
        rows.append({"record_type": "INVOICE" if kind == "purchase_invoices" else "SALES_INVOICE",
                     "original_value": {"sheet": "发票基础信息", "row": row},
                     "source_anchor": {"row": row, "region": f"发票基础信息!A{row}:F{row}"},
                     "normalized_value": {"invoice_no": f"I-{row}", "net_amount": net, "tax": tax, "invoice_total": total},
                     "period_check": "PASS" if row == 2 else "PERIOD_EXCEPTION",
                     "field_sources": {field: {"row": row, "region": f"发票基础信息!{col}{row}"}
                                       for field, col in (("net_amount", "F"), ("tax", "D"), ("invoice_total", "E"))}})
    return {"records": rows, "sheets": [{"sheet": "发票基础信息", "status": "EXTRACTED",
            "reference_rows": [{"row": 4, "reason": reason, "values": [label, None, None, ".03", ".33", ".30"]}]}]}


@pytest.mark.parametrize("kind", ["purchase_invoices", "sales_invoices"])
@pytest.mark.parametrize("reason,label", [("SUMMARY_REFERENCE", "合计"), ("INVOICE_TOTAL", "合计行")])
def test_invoice_three_totals_use_source_columns_and_all_periods(kind, reason, label):
    data = invoice_extraction(kind, reason, label)
    original = deepcopy(data)
    result = reconcile(data, kind)
    assert result["overall"] == "MATCH"
    assert len(result["items"]) == 3
    assert next(i for i in result["items"] if i["code"] == "INVOICE_TOTAL_NET_AMOUNT")["source_regions"][0] == "发票基础信息!F4"
    assert data == original


@pytest.mark.parametrize("column", [3, 4, 5])
def test_invoice_real_cent_difference_requires_parse_review(column):
    data = invoice_extraction()
    data["sheets"][0]["reference_rows"][0]["values"][column] = ".99"
    assert reconcile(data, "purchase_invoices")["overall"] == "REVIEW"


@pytest.mark.parametrize("raw,expected", [(0.1 + 0.2, "MATCH"), ("0.30000000000000004", "MISMATCH"),
                                        (0.301, "MISMATCH"), (0.31, "MISMATCH")])
def test_invoice_float_representation_is_explicit_and_never_cent_tolerance(raw, expected):
    data = invoice_extraction()
    data["sheets"][0]["reference_rows"][0]["values"][5] = raw
    item = next(i for i in reconcile(data, "purchase_invoices")["items"] if i["code"] == "INVOICE_TOTAL_NET_AMOUNT")
    assert item["status"] == expected
    assert item["declared"] == str(raw)
    assert (item.get("comparison") == "FLOAT_REPRESENTATION_ONLY") == (expected == "MATCH")


@pytest.mark.parametrize("problem", ["missing", "subtotal", "multiple", "derived", "drift", "formula"])
def test_invoice_missing_or_ambiguous_anchor_does_not_invent_attestation(problem):
    data = invoice_extraction()
    if problem == "missing": data["sheets"][0]["reference_rows"] = []
    if problem == "subtotal": data["sheets"][0]["reference_rows"][0]["values"][0] = "本页合计"
    if problem == "multiple": data["sheets"][0]["reference_rows"] *= 2
    if problem == "derived": data["records"][0]["field_sources"]["invoice_total"]["region"] = "发票基础信息!D2:F2"
    if problem == "drift": data["records"][0]["field_sources"]["net_amount"]["region"] = "发票基础信息!G2"
    if problem == "formula": data["records"][0]["field_sources"]["tax"]["formula"] = "=F2*0.1"
    assert reconcile(data, "purchase_invoices")["overall"] == "UNATTESTED"


def test_invoice_reference_only_summary_is_not_double_counted():
    data = invoice_extraction()
    summary = deepcopy(data["sheets"][0])
    summary.update(sheet="信息汇总表", status="REFERENCE_ONLY")
    data["sheets"].append(summary)
    assert reconcile(data, "purchase_invoices")["overall"] == "MATCH"
    duplicate = deepcopy(data["records"][0])
    duplicate["original_value"]["sheet"] = "信息汇总表"
    data["records"].append(duplicate)
    assert reconcile(data, "purchase_invoices")["overall"] == "REVIEW"


def test_invoice_signed_red_totals_are_not_absolutized():
    data = invoice_extraction()
    for record in data["records"]:
        for field in ("net_amount", "tax", "invoice_total"):
            record["normalized_value"][field] = "-" + record["normalized_value"][field]
    for index in (3, 4, 5):
        data["sheets"][0]["reference_rows"][0]["values"][index] = "-" + data["sheets"][0]["reference_rows"][0]["values"][index]
    assert reconcile(data, "purchase_invoices")["overall"] == "MATCH"


@pytest.mark.parametrize("label", ["合计", "合计行", "总计"])
def test_tabular_invoice_summary_checks_do_not_change_original_records(label):
    from io import BytesIO
    import openpyxl
    from app.tabular import ParseOptions, extract_workbook

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "发票基础信息"
    sheet.append(["序号", "发票号码", "开票日期", "税额", "价税合计", "金额"])
    sheet.append([1, "I-1", "2026-01-01", "0.01", "0.11", "0.10"])
    sheet.append([2, "I-2", "2026-02-01", "0.02", "0.22", "0.20"])
    sheet.append([label, None, None, "0.03", "0.33", "0.30"])
    # This duplicate presentation must not contribute to the source totals.
    book.copy_worksheet(sheet).title = "信息汇总表"
    content = BytesIO()
    book.save(content)
    book.close()
    result = extract_workbook(content.getvalue(), ParseOptions(document_kind="sales_invoices"), "2026-01")
    assert result["checks"]["overall"] == "MATCH"
    assert len(result["records"]) == 2
    assert result["records"][1]["period_check"] == "PERIOD_EXCEPTION"
    assert result["records"][0]["normalized_value"]["tax_rate"] is None


@pytest.mark.parametrize("kind", ["bank_statement", "purchase_invoices", "sales_invoices"])
def test_tabular_keeps_uncached_formula_only_data_row_for_field_validation(kind):
    from io import BytesIO
    import openpyxl
    from app.tabular import ParseOptions, extract_workbook, read_workbook

    headers = (["交易时间", "收入金额", "支出金额", "账户余额"] if kind == "bank_statement" else
               ["发票号码", "开票日期", "金额", "税额"])
    valid = (["2026-01-01", 10, 0, 110] if kind == "bank_statement" else
             ["I-1", "2026-01-01", 10, 1])
    book = openpyxl.Workbook()
    page = book.active
    page.title = "合成数据"
    page.append(headers)
    page.append(["=1+1"] * len(headers))  # openpyxl supplies no cached values.
    page.append([None] * len(headers))  # A truly empty row must still be ignored.
    page.append(valid)
    content = BytesIO()
    book.save(content)
    book.close()
    raw = read_workbook(content.getvalue())[0]["rows"][1]
    assert raw["values"] == [None] * len(headers) and len(raw["formulas"]) == len(headers)
    result = extract_workbook(content.getvalue(), ParseOptions(document_kind=kind), "2026-01")
    assert [r["source_anchor"]["row"] for r in result["records"]] == [2, 4]
    invalid = result["records"][0]
    assert invalid["original_value"]["values"] == raw["values"]
    assert invalid["original_value"]["formulas"] == raw["formulas"]
    assert any("公式" in issue for issue in invalid["extraction_issues"])
    assert any("日期缺失" in issue for issue in invalid["extraction_issues"])
    assert invalid["extraction_confidence"] == 0.0
    assert result["sheets"][0]["rows"] == 2
