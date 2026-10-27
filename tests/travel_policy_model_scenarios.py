"""Opt-in real-executor acceptance corpus; entirely synthetic, never written to a business ledger.

Run with: python tests/travel_policy_model_scenarios.py --output <new-directory>
This invokes the configured Claude Code provider and consumes model usage.
"""
import argparse
from dataclasses import replace
import json
from pathlib import Path
import sys

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app.autonomous_runtime import run
from app.config import Settings
from app.review_loop import REVIEW_POLICY_VERSION
from app.stage23_pilot import sha256
from app.travel_policy import POLICY_RULES, POLICY_SUMMARY


def main(output):
    output.mkdir(parents=True, exist_ok=False)
    book = openpyxl.Workbook(); bank = book.active; bank.title = '银行回单'
    bank.append(['日期','场景','备注','收入','支出','收款人','报销人','业务说明'])
    rows = [
        ['2026-03-05','普通差旅','差旅费',0,100,'李甲','张乙','本月住宿，收款人为授权代收人；授权在报销单中已列明'],
        ['2026-03-06','混合报销','差旅费',0,400,'张乙','张乙','本月住宿100，餐费100未另述用途，礼品100未另述用途，明确送给客户的定制礼品100'],
        ['2026-03-07','前期已记','差旅费',0,200,'王甲','王甲','结算2月未付住宿费200，前期已挂其他应付款王甲200'],
        ['2026-03-08','前期未记','差旅费',0,300,'刘甲','刘甲','2月业务，2月账内未记，本月首次交票；跨月补记方式未另确认'],
        ['2026-03-09','跨年','差旅费',0,400,'赵甲','赵甲','2025年12月住宿，前期未记；跨年更正方式未确认'],
        ['2026-03-10','补贴999','差旅补贴',0,999,'钱甲','钱甲','金额有付款依据，尚无企业确认标准或天数'],
        ['2026-03-11','补贴1000','差旅补贴',0,1000,'钱乙','钱乙','金额有付款依据，尚无企业确认标准或天数'],
        ['2026-03-12','补贴1001','差旅补贴',0,1001,'钱丙','钱丙','金额有付款依据，尚无企业确认标准或天数'],
        ['2026-03-13','重复报销','差旅费',0,500,'孙甲','孙甲','发票号码INV-DUP，与1月孙甲已报销且已支付的同一发票相同'],
        ['2026-03-14','历史借款','差旅报销冲借款',0,600,'周甲','周甲','报销单列住宿1100、拟冲借款500、实付600，历史借款余额500存在，但本次核销未经人工确认'],
    ]
    for row in rows: bank.append(row)
    invoices = book.create_sheet('发票')
    invoices.append(['日期','场景','发票号码','项目','金额','税额','价税合计','明确用途'])
    for row in [
        ['2026-03-05','普通差旅','INV-1','住宿',100,0,100,'出差住宿'],
        ['2026-03-06','混合报销','INV-2A','住宿',100,0,100,'出差住宿'],
        ['2026-03-06','混合报销','INV-2B','餐费',100,0,100,'未另述用途'],
        ['2026-03-06','混合报销','INV-2C','礼品',100,0,100,'未另述用途'],
        ['2026-03-06','混合报销','INV-2D','客户定制礼品',100,0,100,'已明确为客户业务招待礼品'],
        ['2026-02-20','前期已记','INV-3','住宿',200,0,200,'2月出差'],
        ['2026-02-21','前期未记','INV-4','住宿',300,0,300,'2月出差'],
        ['2025-12-20','跨年','INV-5','住宿',400,0,400,'2025年出差'],
        ['2026-01-20','重复报销','INV-DUP','住宿',500,0,500,'1月孙甲出差'],
        ['2026-03-14','历史借款','INV-10','住宿',1100,0,1100,'3月周甲出差'],
    ]: invoices.append(row)
    history = book.create_sheet('前期凭证及余额')
    history.append(['凭证日期','凭证号','摘要','借方科目','贷方科目','金额','3月期初该笔未结余额'])
    history.append(['2026-02-20','2月记003','王甲INV-3住宿','660201 管理费用—差旅费','224101 其他应付款—王甲',200,200])
    history.append(['2026-01-20','1月记008','孙甲INV-DUP已报销并支付','660201 管理费用—差旅费','1002 银行存款',500,0])
    history.append(['2026-02-25','2月记010','周甲借款','122101 其他应收款—周甲','1002 银行存款',500,500])
    history.append(['说明','本表为全部相关历史','刘甲INV-4及赵甲INV-5前期账内未记；不存在对应应付余额；其他本月发票未曾报销。'])
    accounts = book.create_sheet('企业科目表')
    accounts.append(['编码','名称'])
    for code,name in [('1002','银行存款'),('660201','管理费用—差旅费'),('660202','管理费用—业务招待费'),('660203','管理费用—福利费'),('224101','其他应付款—王甲'),('122101','其他应收款—周甲')]:
        accounts.append([code,name])
    source = output/'synthetic.xlsx'; book.save(source)
    manifest = dict(period='2026-03',sources=[dict(id='synthetic',filename=source.name,path='sources/synthetic.xlsx',sha256=sha256(source))])
    policy = dict(id='synthetic-policy',version=1,enabled=True,summary=POLICY_SUMMARY,rules=POLICY_RULES,
                  source_note='模拟企业已确认本默认口径；没有额外跨月或跨年补记政策',confirmed_by='fixture',confirmed_at='2026-03-31',effective_from='2026-03')
    context = dict(travel_policy=policy,review_loop_version=REVIEW_POLICY_VERSION,statements=[],materials_complete=True,
                   goal='以给定模拟原件独立生成全部业务的有据草稿及具体人工问题，保留场景名称便于追踪。')
    settings = replace(Settings.from_env(),autonomous_backend='claude_code',autonomous_max_rounds=3)
    progress = lambda phase,round_no,message: print(json.dumps(dict(phase=phase,round=round_no,message=message),ensure_ascii=False),flush=True)
    result = run(output/'run',manifest,{'synthetic':source},context,settings,lambda *args:None,progress)
    (output/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print('MODEL_RESULT_SAVED',flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--output',required=True,type=Path)
    main(parser.parse_args().output)
