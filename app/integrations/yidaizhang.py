"""Explicit-template adapter for 易代账 file import and receipt files.

The real vendor workbook is deliberately not inferred here.  A profile must
declare the sheet, header row and column mapping after the customer supplies
the actual template.  This prevents a plausible-looking but wrong voucher
file from being produced.
"""
from __future__ import annotations

import csv
import hashlib
import json
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.ontology.store import digest


ADAPTER_VERSION = "yidaizhang-file-v1"
RECEIPT_STATUSES = {"IMPORTED", "FAILED", "PARTIAL"}
REQUIRED_VOUCHER_FIELDS = {
    "voucher_number", "voucher_date", "summary", "account_name", "debit", "credit",
}


class YidaizhangFormatError(ValueError):
    """The supplied template/profile/receipt does not satisfy the adapter contract."""


class YidaizhangBindingError(ValueError):
    """A file does not bind to the exact FinWise export it claims to acknowledge."""


@dataclass(frozen=True)
class TemplateProfile:
    sheet: str
    header_row: int
    columns: dict[str, str]
    receipt_status_values: dict[str, str] | None = None

    @classmethod
    def from_dict(cls, value: dict[str, Any]) -> "TemplateProfile":
        if not isinstance(value, dict) or not isinstance(value.get("sheet"), str) or not value["sheet"].strip():
            raise YidaizhangFormatError("易代账模板配置必须明确 sheet")
        if type(value.get("header_row")) is not int or value["header_row"] < 1:
            raise YidaizhangFormatError("易代账模板配置必须明确正整数 header_row")
        columns = value.get("columns")
        if not isinstance(columns, dict):
            raise YidaizhangFormatError("易代账模板配置必须明确 columns 映射")
        missing = sorted(REQUIRED_VOUCHER_FIELDS - set(columns))
        if missing:
            raise YidaizhangFormatError(f"易代账模板配置缺少字段映射: {', '.join(missing)}")
        if any(not isinstance(key, str) or not key.strip() for key in columns.values()):
            raise YidaizhangFormatError("易代账模板列名不能为空")
        statuses = value.get("receipt_status_values")
        if statuses is not None and (not isinstance(statuses, dict) or any(not isinstance(k, str) or not isinstance(v, str) for k, v in statuses.items())):
            raise YidaizhangFormatError("receipt_status_values 必须是字符串映射")
        return cls(sheet=value["sheet"], header_row=value["header_row"], columns=dict(columns), receipt_status_values=statuses)


def file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def inspect_template(path: Path) -> dict[str, Any]:
    """Inspect workbook structure without guessing business columns."""
    path = Path(path)
    if not path.is_file():
        raise YidaizhangFormatError(f"模板文件不存在: {path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise YidaizhangFormatError(f"无法读取模板文件: {exc}") from exc
    sheets = []
    for worksheet in workbook.worksheets:
        rows = []
        for row in worksheet.iter_rows(min_row=1, max_row=min(8, worksheet.max_row), values_only=True):
            rows.append([None if value is None else str(value) for value in row[:30]])
        sheets.append({"title": worksheet.title, "max_row": worksheet.max_row, "max_column": worksheet.max_column, "preview": rows})
    return {"adapter_version": ADAPTER_VERSION, "filename": path.name, "sha256": file_sha256(path), "sheets": sheets}


def export_voucher_file(
    *,
    export: dict[str, Any],
    scope: dict[str, Any],
    template_path: Path,
    profile: TemplateProfile,
    output_path: Path,
    manifest_path: Path,
) -> dict[str, Any]:
    """Create a workbook using only an explicit customer-supplied mapping."""
    template_info = inspect_template(template_path)
    if profile.sheet not in {item["title"] for item in template_info["sheets"]}:
        raise YidaizhangFormatError(f"模板不存在配置的工作表: {profile.sheet}")
    lines = export.get("data", export).get("lines")
    if not isinstance(lines, list) or not lines:
        raise YidaizhangFormatError("导出对象没有凭证分录")
    export_data = export.get("data", export)
    package_id = export_data.get("package_id")
    voucher_version_id = export_data.get("voucher_version_id")
    voucher_version = export_data.get("voucher_version")
    if not package_id or not voucher_version_id or type(voucher_version) is not int:
        raise YidaizhangBindingError("导出对象缺少 package_id、voucher_version_id 或 voucher_version")
    header = export_data.get("voucher_header")
    if not isinstance(header, dict) or any(not header.get(key) for key in ("voucher_number", "voucher_date", "summary")):
        raise YidaizhangBindingError("导出对象缺少人工确认的 voucher_header（凭证号、日期、摘要）")
    rows = _voucher_rows(lines, profile, header)
    workbook = load_workbook(template_path)
    worksheet = workbook[profile.sheet]
    header_values = [cell.value for cell in worksheet[profile.header_row]]
    header_lookup = {str(value): index + 1 for index, value in enumerate(header_values) if value is not None}
    for field, column_name in profile.columns.items():
        if column_name not in header_lookup:
            raise YidaizhangFormatError(f"模板工作表缺少配置列: {column_name}")
    existing_rows = list(worksheet.iter_rows(min_row=profile.header_row + 1, values_only=True))
    if any(any(value is not None and str(value).strip() for value in row) for row in existing_rows):
        raise YidaizhangFormatError("导入模板在标题行后已有数据，拒绝覆盖以免重复导入；请提供空白模板")
    start_row = profile.header_row + 1
    for row_index, row in enumerate(rows, start=start_row):
        for field, column_name in profile.columns.items():
            worksheet.cell(row=row_index, column=header_lookup[column_name], value=row.get(field))
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    manifest = {
        "adapter_version": ADAPTER_VERSION,
        "scope": scope,
        "period": scope.get("accounting_period_id"),
        "package_id": package_id,
        "export_id": export.get("object_id") or export.get("data", export).get("export_id"),
        "voucher_version_id": voucher_version_id,
        "voucher_version": voucher_version,
        "template": {"filename": template_path.name, "sha256": template_info["sha256"], "profile": {"sheet": profile.sheet, "header_row": profile.header_row, "columns": profile.columns}},
        "file": {"filename": output_path.name, "sha256": file_sha256(output_path)},
    }
    manifest_path = Path(manifest_path)
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    return manifest


def _voucher_rows(lines: list[dict[str, Any]], profile: TemplateProfile, header: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for index, line in enumerate(lines, start=1):
        if not isinstance(line, dict):
            raise YidaizhangFormatError(f"第 {index} 条分录不是对象")
        account = line.get("account")
        if not isinstance(account, str) or not account.strip():
            raise YidaizhangFormatError(f"第 {index} 条分录缺少科目")
        amount = _amount(line.get("amount"), index)
        direction = line.get("direction")
        if direction not in {"DEBIT", "CREDIT"}:
            raise YidaizhangFormatError(f"第 {index} 条分录借贷方向非法")
        row = {
            "voucher_number": line.get("voucher_number") or header["voucher_number"],
            "voucher_date": _date_value(line.get("voucher_date") or line.get("date") or header["voucher_date"]),
            "summary": line.get("summary") or header["summary"],
            "account_name": account,
            "debit": amount if direction == "DEBIT" else None,
            "credit": amount if direction == "CREDIT" else None,
        }
        if "account_code" in profile.columns:
            code = line.get("account_code")
            if not isinstance(code, str) or not code.strip():
                raise YidaizhangFormatError(f"第 {index} 条分录缺少模板要求的科目编码")
            row["account_code"] = code
        if "line_no" in profile.columns:
            row["line_no"] = line.get("line_no", index)
        rows.append(row)
    return rows


def _amount(value: Any, index: int) -> float:
    try:
        amount = Decimal(str(value))
        if not amount.is_finite() or amount < 0 or amount != amount.quantize(Decimal("0.01")):
            raise InvalidOperation
    except (InvalidOperation, ValueError):
        raise YidaizhangFormatError(f"第 {index} 条分录金额不是精确到分的有限数")
    return float(amount)


def _date_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value
    if isinstance(value, str) and value.strip():
        return value.strip()
    raise YidaizhangFormatError("凭证日期必须是日期或非空字符串")


def load_receipt(path: Path, profile: TemplateProfile | None = None) -> dict[str, Any]:
    """Read JSON directly; tabular receipts require an explicit profile."""
    path = Path(path)
    if not path.is_file():
        raise YidaizhangFormatError(f"回执文件不存在: {path}")
    if path.suffix.lower() == ".json":
        try:
            value = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise YidaizhangFormatError(f"回执 JSON 无法读取: {exc}") from exc
        if not isinstance(value, dict):
            raise YidaizhangFormatError("回执 JSON 必须是对象")
        payload = dict(value)
    elif path.suffix.lower() in {".csv", ".xlsx", ".xlsm"}:
        if profile is None:
            raise YidaizhangFormatError("表格回执必须先提供明确列映射，不能猜测字段")
        payload = _load_tabular_receipt(path, profile)
    else:
        raise YidaizhangFormatError("回执仅支持 JSON、CSV 或 XLSX；未知格式必须先人工转换")
    payload["receipt_file_name"] = path.name
    payload["receipt_file_sha256"] = file_sha256(path)
    payload["adapter_version"] = ADAPTER_VERSION
    return validate_receipt_payload(payload, profile=profile)


def _load_tabular_receipt(path: Path, profile: TemplateProfile) -> dict[str, Any]:
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        if len(rows) != 1:
            raise YidaizhangFormatError("表格回执必须明确只有一条批次回执")
        source = rows[0]
    else:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if profile.sheet not in workbook.sheetnames:
            raise YidaizhangFormatError(f"回执工作表不存在: {profile.sheet}")
        worksheet = workbook[profile.sheet]
        values = list(worksheet.iter_rows(min_row=profile.header_row, max_row=profile.header_row + 1, values_only=True))
        if len(values) != 2:
            raise YidaizhangFormatError("表格回执必须包含标题行和一条数据")
        headers = [str(item) for item in values[0]]
        source = dict(zip(headers, values[1]))
    def get(field: str) -> Any:
        column = profile.columns.get(field)
        return source.get(column) if column else None
    return {key: get(key) for key in ("status", "external_batch", "export_id", "voucher_version_id", "voucher_version", "period", "export_file_sha256", "reason") if get(key) is not None}


def validate_receipt_payload(payload: dict[str, Any], *, profile: TemplateProfile | None = None) -> dict[str, Any]:
    if not isinstance(payload, dict):
        raise YidaizhangFormatError("回执必须是对象")
    status = payload.get("status")
    if profile and profile.receipt_status_values:
        reverse = {value: key for key, value in profile.receipt_status_values.items()}
        status = reverse.get(status, status)
    if status not in RECEIPT_STATUSES:
        raise YidaizhangFormatError("回执状态必须是 IMPORTED、FAILED 或 PARTIAL，未知状态不得自动映射")
    external_batch = payload.get("external_batch")
    if not isinstance(external_batch, str) or not external_batch.strip():
        raise YidaizhangFormatError("回执缺少 external_batch")
    for key in ("export_id", "voucher_version_id"):
        if not isinstance(payload.get(key), str) or not payload[key].strip():
            raise YidaizhangFormatError(f"回执缺少 {key}")
    if type(payload.get("voucher_version")) is not int or payload["voucher_version"] < 1:
        raise YidaizhangFormatError("回执 voucher_version 必须是正整数")
    normalized = dict(payload)
    normalized["adapter"] = ADAPTER_VERSION
    normalized["status"] = status
    normalized["external_batch"] = external_batch.strip()
    normalized["payload_hash"] = digest({key: normalized.get(key) for key in ("status", "external_batch", "export_id", "voucher_version_id", "voucher_version", "reason")})
    return normalized


def bind_receipt_to_manifest(payload: dict[str, Any], manifest: dict[str, Any]) -> dict[str, Any]:
    expected = {
        "export_id": manifest.get("export_id"),
        "voucher_version_id": manifest.get("voucher_version_id"),
        "voucher_version": manifest.get("voucher_version"),
    }
    if any(payload.get(key) != value for key, value in expected.items()):
        raise YidaizhangBindingError("易代账回执与导出包、凭证版本或修订号不一致")
    if payload.get("period") != manifest.get("period"):
        raise YidaizhangBindingError("易代账回执期间与导出 Manifest 不一致或缺失")
    if payload.get("export_file_sha256") != manifest.get("file", {}).get("sha256"):
        raise YidaizhangBindingError("易代账回执没有匹配导出文件哈希，不能自动绑定")
    return {**payload, "adapter": ADAPTER_VERSION, "manifest_file_sha256": digest(manifest)}
