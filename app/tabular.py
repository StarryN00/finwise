"""Local deterministic spreadsheet extraction. Never executes formulas or model output."""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
import re
from typing import Literal, Optional
from zipfile import BadZipFile, ZipFile
from xml.etree import ElementTree

from pydantic import BaseModel, ConfigDict, Field, StrictStr

PARSER_VERSION = "tabular-v5"
MAX_BYTES = 16 * 1024 * 1024
MAX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_CELLS = 200_000
MAX_ROWS = 20_000
MAX_COLUMNS = 256


class ParseOptions(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)
    document_kind: Literal[
        "purchase_invoices", "sales_invoices", "bank_statement", "payroll",
        "social_security", "housing_fund", "electronic_acceptance", "individual_income_tax",
        "opening_balance", "contract", "stock_in",
    ]
    bank_account_ref: Optional[StrictStr] = Field(default=None, min_length=1, max_length=128)


class ExtractionError(ValueError):
    pass


def column_name(index):
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def serializable(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def read_workbook(content):
    if len(content) > MAX_BYTES:
        raise ExtractionError("文件超过解析大小限制，请拆分资料")
    sheets, cell_count = [], 0
    try:
        if content.startswith(b"PK"):
            import openpyxl
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 1000 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
                    raise ExtractionError("工作簿展开大小超过限制")
                for entry in entries:
                    if entry.filename.endswith(".xml"):
                        xml = archive.read(entry).upper()
                        if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                            raise ExtractionError("不接受包含外部实体声明的工作簿")
            raw = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            cached = None
            try:
                cached = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
                if len(raw.worksheets) > 50:
                    raise ExtractionError("工作表数量超过限制")
                for sheet, value_sheet in zip(raw.worksheets, cached.worksheets):
                    sheet.reset_dimensions()
                    value_sheet.reset_dimensions()
                    rows = []
                    for row_no, (cells, values) in enumerate(zip(sheet.iter_rows(), value_sheet.iter_rows()), 1):
                        cell_count += len(cells)
                        if row_no > MAX_ROWS or len(cells) > MAX_COLUMNS or cell_count > MAX_CELLS:
                            raise ExtractionError("工作表行列或单元格数量超过限制")
                        rows.append({"row": row_no, "values": [serializable(v.value) for v in values],
                                     "types": [c.data_type for c in cells],
                                     "number_formats": [getattr(c, "number_format", None) for c in cells],
                                     "formulas": {c.coordinate: c.value for c in cells if c.data_type == "f"}})
                    with ZipFile(BytesIO(content)) as archive:
                        xml = ElementTree.fromstring(archive.read(sheet._worksheet_path))
                        merges = [openpyxl.utils.range_boundaries(node.attrib["ref"])
                                  for node in xml.findall("{*}mergeCells/{*}mergeCell")]
                        hidden_rows = sorted({int(node.attrib["r"]) for node in xml.findall("{*}sheetData/{*}row")
                                              if node.get("hidden", "").strip().lower() in {"1", "true"}})
                        hidden_columns = set()
                        for node in xml.findall("{*}cols/{*}col"):
                            if node.get("hidden", "").strip().lower() in {"1", "true"}:
                                first, last = int(node.attrib["min"]), int(node.attrib["max"])
                                if not 1 <= first <= last <= 16384:
                                    raise ExtractionError("工作簿隐藏列范围无效")
                                hidden_columns.update(range(first, last + 1))
                    sheets.append({"name": sheet.title, "rows": rows, "merged_cells": merges,
                                   "hidden_rows": hidden_rows, "hidden_columns": sorted(hidden_columns),
                                   "sheet_state": sheet.sheet_state})
            finally:
                raw.close()
                if cached is not None:
                    cached.close()
        elif content.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
            import xlrd
            try:
                book = xlrd.open_workbook(file_contents=content, on_demand=True)
            except (xlrd.compdoc.CompDocError, xlrd.biffh.XLRDError, ValueError, OSError, TypeError) as exc:
                raise ExtractionError("XLS 工作簿损坏或无法解析，请检查原件完整性") from exc
            layout_book = None
            try:
                if book.nsheets > 50:
                    raise ExtractionError("工作表数量超过限制")
                # Formatting-only trailing cells must not change persisted row
                # anchors. Read values at the legacy width and merges separately.
                try:
                    layout_book = xlrd.open_workbook(file_contents=content, on_demand=True, formatting_info=True)
                except (xlrd.compdoc.CompDocError, xlrd.biffh.XLRDError) as exc:
                    raise ExtractionError("XLS 合并表头信息损坏，请检查原件完整性") from exc
                for sheet_index, sheet in enumerate(book.sheets()):
                    layout_sheet = layout_book.sheet_by_index(sheet_index)
                    cell_count += sheet.nrows * sheet.ncols
                    if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLUMNS or cell_count > MAX_CELLS:
                        raise ExtractionError("工作表行列或单元格数量超过限制")
                    rows = []
                    for r in range(sheet.nrows):
                        values, types = [], []
                        for cell in sheet.row(r):
                            value = cell.value
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                value = xlrd.xldate.xldate_as_datetime(value, book.datemode).isoformat()
                            elif cell.ctype == xlrd.XL_CELL_ERROR:
                                # Keep the error marker as source text; never coerce it to a number.
                                value = f"#ERROR:{value}"
                            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                                value = bool(value)
                            values.append(value)
                            types.append(str(cell.ctype))
                        number_formats = []
                        for column in range(len(values)):
                            # Some older readers/test doubles have no formatting
                            # table. Preserve unknown rather than inventing a format.
                            if not hasattr(layout_book, "xf_list"):
                                number_formats.append(None)
                                continue
                            xf = layout_book.xf_list[layout_sheet.cell_xf_index(r, column)]
                            cell_format = layout_book.format_map.get(xf.format_key)
                            number_formats.append(cell_format.format_str if cell_format else None)
                        rows.append({"row": r+1, "values": values, "types": types, "formulas": {},
                                     "number_formats": number_formats})
                    sheets.append({"name": sheet.name, "rows": rows,
                                   "merged_cells": [(c1+1, r1+1, c2, r2) for r1, r2, c1, c2 in layout_sheet.merged_cells],
                                   "hidden_rows": sorted(r+1 for r, info in getattr(layout_sheet, "rowinfo_map", {}).items() if info.hidden),
                                   "hidden_columns": sorted(c+1 for c, info in getattr(layout_sheet, "colinfo_map", {}).items() if info.hidden),
                                   "sheet_state": {0: "visible", 1: "hidden", 2: "veryHidden"}.get(getattr(layout_sheet, "visibility", 0), "unknown")})
            finally:
                book.release_resources()
                if layout_book is not None:
                    layout_book.release_resources()
        else:
            raise ExtractionError("无法识别 Excel 格式，暂支持真实 XLS 和 XLSX 文件")
    except ExtractionError:
        raise
    except (ImportError, BadZipFile, ValueError, KeyError, OSError, TypeError) as exc:
        raise ExtractionError("工作簿无法解析，请检查文件完整性及解析依赖") from exc
    return sheets


INVOICE_FIELDS = {
    "invoice_no": ("数电发票号码", "发票号码"), "invoice_date": ("开票日期",),
    "net_amount": ("金额", "不含税金额"), "tax": ("税额",), "tax_rate": ("税率",), "invoice_total": ("价税合计",),
    "seller_name": ("销售方纳税人名称", "销方名称", "销售方名称"),
    "seller_tax_id": ("销售方纳税人识别号", "销方识别号"),
    "buyer_name": ("购买方名称", "购方名称"), "buyer_tax_id": ("购方识别号", "购买方纳税人识别号"),
    "invoice_status": ("发票状态",),
}
BANK_FIELDS = {
    "transaction_date": ("交易日期", "交易时间"), "expense": ("支出金额", "支出", "借方金额"),
    "income": ("收入金额", "收入", "贷方金额"), "balance": ("账户余额", "余额"),
    "counterparty": ("对方户名", "对手方名称"), "counterparty_account": ("对方账号",),
    "transaction_id": ("交易流水号", "流水号"), "summary": ("摘要",), "note": ("附言", "用途"),
}
GENERIC_FIELDS = {
    "payroll": {
        "person_name": ("姓名",), "basic_salary": ("基本工资",), "actual_salary": ("实发工资",),
        "tax": ("个税",), "employer_social": ("公司社保", "单位社保"), "employee_social": ("个人社保",),
        "employer_housing_fund": ("单位公积金", "公司公积金"), "employee_housing_fund": ("个人公积金",),
        "housing_fund": ("公积金", "公积金公司/个人"),
    },
    "social_security": {
        "person_id": ("个人编号",), "person_name": ("姓名",), "period_ref": ("对应费款所属期", "结算期"),
        "insurance_type": ("险种类型",), "base": ("缴费基数总额",), "employer_amount": ("单位缴费金额",),
        "employee_amount": ("个人缴费金额",), "total": ("合计缴费金额",), "payment_status": ("缴费标志",),
        "arrival_date": ("到账时间",),
    },
    "housing_fund": {
        "account": ("个人公积金账号",), "person_name": ("姓名",), "business_type": ("业务类别",),
        "amount": ("缴存金额",), "subsidy": ("其中补贴金额",), "start_period": ("缴存起始月份",),
        "end_period": ("缴存终止月份",), "entry_date": ("入账日期",),
    },
    "electronic_acceptance": {
        "acceptance_no": ("电子票据号", "票据包号"), "sub_range": ("子票区间",),
        "counterparty": ("对手方名称", "前手信息"), "acceptance_type": ("票据类型",),
        "issue_date": ("出票日期",), "maturity_date": ("到期日期", "票面到期日期"),
        "amount": ("票据金额", "票面金额", "交易金额"), "transaction_type": ("交易种类", "业务类型"),
        "transaction_date": ("交易时间", "自动任务发起日"), "status": ("交易状态", "交易结果"),
    },
    "individual_income_tax": {
        "person_name": ("姓名",), "person_id": ("身份证件号码",), "income_item": ("所得项目",),
        "income": ("收入",), "expense": ("费用",), "tax_free_income": ("免税收入",),
        "pension": ("基本养老保险费",), "medical": ("基本医疗保险费",), "unemployment": ("失业保险费",),
        "housing_fund": ("住房公积金",),
    },
    "opening_balance": {
        "account_code": ("科目编码", "科目代码", "科目编号", "会计科目编码"),
        "account_name": ("科目名称", "会计科目", "科目"),
        "closing_debit": ("期末借方", "期末借方余额", "期末借方金额", "借方余额"),
        "closing_credit": ("期末贷方", "期末贷方余额", "期末贷方金额", "贷方余额"),
        "opening_debit": ("期初借方", "期初借方余额", "期初借方金额"),
        "opening_credit": ("期初贷方", "期初贷方余额", "期初贷方金额"),
        "requires_auxiliary": ("是否辅助核算", "辅助核算", "含辅助核算"),
        "auxiliary_key": ("辅助核算项", "辅助项", "辅助核算编码", "辅助核算名称"),
    },
    "contract": {
        "contract_no": ("合同编号", "合同号", "协议编号"), "contract_date": ("合同日期", "签订日期", "签署日期"),
        "supplier": ("供应商名称", "供应商", "乙方名称"), "amount": ("合同金额", "含税合同金额", "金额"),
    },
    "stock_in": {
        "stock_in_no": ("入库单号", "入库编号", "收货单号"), "stock_in_date": ("入库日期", "入库时间", "收货日期"),
        "supplier": ("供应商名称", "供应商", "供货单位"), "amount": ("入库金额", "入库含税金额", "金额"),
        "item": ("存货名称", "物料名称", "商品名称"),
    },
}
GENERIC_RECORD_TYPES = {
    "payroll": "PAYROLL",
    "social_security": "SOCIAL_SECURITY",
    "housing_fund": "HOUSING_FUND",
    "electronic_acceptance": "ELECTRONIC_ACCEPTANCE",
    "individual_income_tax": "INDIVIDUAL_INCOME_TAX",
    "opening_balance": "OPENING_BALANCE",
    "contract": "CONTRACT",
    "stock_in": "STOCK_IN",
}


def header_mapping(row, kind):
    labels = [re.sub(r"\s+", "", str(x or "")) for x in row]
    fields = BANK_FIELDS if kind == "bank_statement" else INVOICE_FIELDS if kind in {"purchase_invoices", "sales_invoices"} else GENERIC_FIELDS[kind]
    mapping = {field: [labels.index(name) for name in names if name in labels] for field, names in fields.items()}
    if kind == "bank_statement":
        required = ("transaction_date", "expense", "income")
    elif kind in {"purchase_invoices", "sales_invoices"}:
        required = ("invoice_no", "invoice_date", "net_amount", "tax")
    else:
        required = {
            "payroll": ("person_name", "actual_salary"),
            "social_security": ("person_id", "period_ref", "insurance_type", "employer_amount", "employee_amount"),
            "housing_fund": ("account", "amount", "entry_date"),
            "electronic_acceptance": ("acceptance_no", "amount", "transaction_date"),
            "individual_income_tax": ("person_name", "income"),
            "opening_balance": ("account_code", "account_name", "closing_debit", "closing_credit", "opening_debit", "opening_credit", "requires_auxiliary"),
            "contract": ("contract_no", "contract_date", "supplier", "amount"),
            "stock_in": ("stock_in_no", "stock_in_date", "supplier", "amount"),
        }[kind]
        # The individual income tax template places the nested "收入" header
        # several rows below "姓名". Its stable column is H even when the
        # merged header is not materialized by xlrd/openpyxl.
        if kind == "individual_income_tax" and mapping["person_name"] and not mapping["income"]:
            mapping["income"] = [7]
    return mapping if all(mapping[field] for field in required) else None


def money_value(value):
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        raise ValueError("金额缺失")
    text = str(value).strip()
    if "," in text:
        if not re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", text):
            raise ValueError("金额分组格式无效")
        text = text.replace(",", "")
    try:
        value = Decimal(text)
        if not value.is_finite() or abs(value) >= Decimal("1e12") or value != value.quantize(Decimal(".01")):
            raise ValueError("金额非有限、超限或精度超过分")
        return format(value, ".2f")
    except InvalidOperation as exc:
        raise ValueError("金额格式无效") from exc


def date_value(value):
    text = str(value or "").strip()
    match = re.match(r"^(\d{4})[-/.年]?(\d{2})[-/.月]?(\d{2})(?:日|T|\s|$)", text)
    if not match:
        raise ValueError("日期缺失或格式不能确定")
    return date(*map(int, match.groups())).isoformat()


def period_value(value):
    text = str(value or "").strip()
    match = re.search(r"(20\d{2})\D?((?:0[1-9])|(?:1[0-2]))(?:\D|$)", text)
    if not match:
        match = re.search(r"(20\d{2})(0[1-9]|1[0-2])", text)
    if not match:
        raise ValueError("期间缺失或格式不能确定")
    return f"{match.group(1)}-{int(match.group(2)):02d}"


def detect_workbook_period(sheets):
    patterns = (
        r"(20\d{2})年\D{0,3}(0?[1-9]|1[0-2])月",
        r"(20\d{2})[.\-/](0?[1-9]|1[0-2])月",
        r"(?<!\d)(20\d{2})(0[1-9]|1[0-2])(?!\d)",
    )
    for sheet in sheets:
        for row in sheet["rows"]:
            for value in row["values"]:
                text = str(value or "")
                for pattern in patterns:
                    match = re.search(pattern, text)
                    if match:
                        return f"{match.group(1)}-{int(match.group(2)):02d}"
    return None


def cell_source(sheet, row, headers, index, header_paths=None):
    coordinate = f"{column_name(index+1)}{row['row']}"
    value = row["values"][index] if index < len(row["values"]) else None
    path = (header_paths or {}).get(index, [])
    label = str(headers[index] or "") if index < len(headers) else ""
    if path:
        label = " / ".join(str(item["raw_value"]) for item in path)
    return {"row": row["row"], "region": f"{sheet}!{coordinate}", "label": label,
            "raw_value": value, "rawvalue": value, "original_value": value,
            "header": label, "source_label": re.sub(r"\s+", " ", label).strip(), "header_path": path,
            "formula": row["formulas"].get(coordinate)}


def payroll_header(sheet, start):
    """Use explicit labels and actual merge bounds, never fill blank groups sideways."""
    rows = sheet["rows"]
    clean = lambda value: re.sub(r"\s+", "", str(value or ""))
    if "姓名" not in [clean(v) for v in rows[start]["values"]]:
        return None
    end = start
    for candidate in rows[start+1:start+4]:
        labels = [clean(v) for v in candidate["values"]]
        if any(isinstance(v, (int, float)) for v in candidate["values"]) or "姓名" in labels:
            break
        if not any(v in {"单位缴", "公司缴", "个人缴", "公司社保", "个人社保", "公积金公司/个人", "个税", "实发工资"}
                   for v in labels):
            break
        end += 1
    width = max(len(r["values"]) for r in rows[start:end+1])
    paths = {i: [] for i in range(width)}
    for row in rows[start:end+1]:
        for index in range(width):
            r, c = row["row"], index+1
            bounds = next((b for b in sheet.get("merged_cells", []) if b[0] <= c <= b[2] and b[1] <= r <= b[3]), (c, r, c, r))
            c1, r1, c2, r2 = bounds
            if not rows[start]["row"] <= r1 <= rows[end]["row"]:
                continue
            source_row = rows[r1-1]
            value = source_row["values"][c1-1] if c1 <= len(source_row["values"]) else None
            if value in (None, ""):
                continue
            region = f"{sheet['name']}!{column_name(c1)}{r1}"
            if (c1, r1) != (c2, r2):
                region += f":{column_name(c2)}{r2}"
            item = {"row": r1, "region": region, "raw_value": value,
                    "formula": source_row["formulas"].get(f"{column_name(c1)}{r1}")}
            if item not in paths[index]:
                paths[index].append(item)
    mapping = {field: [] for field in GENERIC_FIELDS["payroll"]}
    unresolved = {}
    for index, path in paths.items():
        labels = [clean(item["raw_value"]) for item in path]
        joined = "/".join(labels)
        generic_side = any(v in {"单位缴", "公司缴", "个人缴", "单位缴费", "个人缴费"} for v in labels)
        social = any(v in {"社保", "社会保险", "公司社保", "单位社保", "个人社保"} for v in labels)
        housing = any(v in {"公积金", "住房公积金", "公积金公司/个人", "公司公积金", "单位公积金", "个人公积金"} for v in labels)
        unknown_benefit = any(("社保" in v or "社会保险" in v or "公积金" in v) and v not in {
            "社保", "社会保险", "公司社保", "单位社保", "个人社保", "公积金", "住房公积金",
            "公积金公司/个人", "公司公积金", "单位公积金", "个人公积金"} for v in labels)
        # A shared '公司/个人' group describes the fund, not either payer.
        side_labels = [v for v in labels if "公司/个人" not in v]
        employer = any("单位" in v or "公司" in v for v in side_labels)
        employee = any("个人" in v for v in side_labels)
        if any(item["formula"] for item in path):
            unresolved[index] = "表头含公式，须核对字段映射"
        elif unknown_benefit:
            unresolved[index] = f"表头“{joined}”未明确属于已支持的缴费金额列"
        elif (social or housing) and employer != employee and social != housing:
            field = ("employer_" if employer else "employee_") + ("social" if social else "housing_fund")
            mapping[field].append(index)
        elif social and not employer and not employee:
            unresolved[index] = f"表头“{joined}”未明确单位或个人缴费方，不能映射社保金额"
        elif generic_side or (social and housing) or ((social or housing) and employer and employee):
            unresolved[index] = f"表头“{joined}”无法明确区分社保/公积金及缴费方"
        else:
            for field, aliases in GENERIC_FIELDS["payroll"].items():
                if any(label in aliases for label in labels):
                    # Shared group labels must not become a one-sided total.
                    if field == "housing_fund" and (len(labels) > 1 or "公积金公司/个人" in labels):
                        unresolved[index] = f"表头“{joined}”未明确公积金金额口径"
                    else:
                        mapping[field].append(index)
    if not mapping["person_name"] or not mapping["actual_salary"]:
        return None
    duplicates = {field: indices[:] for field, indices in mapping.items() if len(indices) > 1}
    for field, indices in duplicates.items():
        for index in indices:
            unresolved[index] = f"{field}：多个来源列重复映射，须明确取值列"
        mapping[field] = []
    return {"end": end, "mapping": mapping, "header_paths": paths, "unresolved": unresolved,
            "header_rows": [{"row": r["row"], "values": r["values"]} for r in rows[start:end+1]]}


def payroll_period_source(sheet, rows, previous=None):
    candidates = []
    for row in rows:
        for index, value in enumerate(row["values"]):
            text = str(value or "")
            if not re.search(r"月份|所属期|工资期间|工资表|工资条", text):
                continue
            matches = re.findall(r"(?<!\d)(20\d{2})[年./-]?(0?[1-9]|1[0-2])(?:月|(?=\D|$))", text)
            source = cell_source(sheet["name"], row, row["values"], index)
            source.update({"label": "工资业务期间", "source_label": "工资业务期间（工资表标题）",
                           "source_kind": "WORKSHEET_TITLE", "derivation": "从工资表标题提取年月（YYYY-MM），须核对业务归属"})
            for year, month in matches or [(None, None)]:
                candidates.append({**source, "value": f"{year}-{int(month):02d}" if year else None})
    if not candidates:
        return previous or {"label": "工资业务期间", "source_label": "工资业务期间", "original_value": None,
                            "raw_value": None, "rawvalue": None, "row": None, "region": "", "status": "MISSING", "value": None}
    values = {c["value"] for c in candidates}
    if len(values) == 1 and None not in values and not any(c["formula"] for c in candidates):
        return {**candidates[0], "status": "EXPLICIT", "candidates": candidates}
    return {"label": "工资业务期间", "source_label": "工资业务期间", "original_value": None,
            "raw_value": None, "rawvalue": None, "row": None, "region": "",
            "value": None, "status": "AMBIGUOUS", "candidates": candidates}


def extract_payroll_sheet(sheet, period):
    records, context, period_source = [], None, None
    previous_header, header_end = 0, 0
    first_header = None
    for index, row in enumerate(sheet["rows"]):
        if context and index <= header_end:
            continue
        new_header = payroll_header(sheet, index)
        if new_header:
            first_header = first_header or row["row"]
            period_source = payroll_period_source(sheet, sheet["rows"][previous_header:index], period_source)
            context = new_header
            context["period_source"] = period_source
            previous_header, header_end = new_header["end"]+1, new_header["end"]
            continue
        if not context:
            continue
        mapping = context["mapping"]
        # Ambiguous identity/amount headers cannot be used to select employees.
        if not mapping["person_name"] or not mapping["actual_salary"]:
            continue
        person_index, salary_index = mapping["person_name"][0], mapping["actual_salary"][0]
        person = str(row["values"][person_index] or "").strip() if person_index < len(row["values"]) else ""
        salary = row["values"][salary_index] if salary_index < len(row["values"]) else None
        if not person or person in {"合计", "总计", "小计", "本页合计", "公司社保", "个人社保", "社保明细", "社保拆分"} or person.endswith("有限公司"):
            continue
        if person.startswith(("薪资", "说明", "备注", "注：", "注:", "月份:", "月份：", "工资期间", "所属期")):
            continue
        # Dated employment/pay-change notes occupy the name column in wage
        # slips. Exclude only explicit notes with a genuinely empty salary cell.
        if salary in (None, "") and f"{column_name(salary_index+1)}{row['row']}" not in row["formulas"] and re.match(
            r"^(?:20\d{2}|\d{2})[./-](?:0?[1-9]|1[0-2])"
            r"(?:[./-](?:0?[1-9]|[12]\d|3[01]))?月?\s*(?:入职|上班|调薪|涨薪|调\s*\d)", person
        ):
            continue
        # Exclude explicitly labelled/merged notes, not rows merely because an
        # employee's required amount is blank, malformed, or a formula.
        if salary in (None, "") and (
            re.search(r"20\d{2}.*工资[表条]", person)
            or any(c1 <= person_index+1 <= c2 and c1 <= salary_index+1 <= c2 and r1 <= row["row"] <= r2
                   for c1, r1, c2, r2 in sheet.get("merged_cells", []))
        ):
            continue
        headers = context["header_rows"][0]["values"]
        records.append(extract_generic_row(sheet["name"], row, headers, mapping, "payroll", period, None, None, context))
    periods = {record["normalized_value"]["period"] for record in records}
    return {"records": records, "sheets": [{"sheet": sheet["name"], "status": "EXTRACTED" if context else "UNRECOGNIZED", "rows": len(records),
                                            "header_row": first_header, "document_period": next(iter(periods)) if len(periods) == 1 else None}]}


def extract_generic_row(sheet, row, headers, mapping, kind, period, bank_account, document_period, payroll_context=None):
    values, number, issues, fields = row["values"], row["row"], [], {}

    def raw(field):
        for index in mapping.get(field, []):
            value = values[index] if index < len(values) else None
            fields[field] = cell_source(sheet, row, headers, index, (payroll_context or {}).get("header_paths"))
            if column_name(index + 1) + str(number) in row["formulas"]:
                issues.append(field + "：公式值须与原件计算结果核实")
            if value not in (None, ""):
                return value
        return None

    def text(field, required=False, identifier=False):
        value = raw(field)
        if value is None:
            if required:
                issues.append(field + "：来源缺失")
            return None
        if identifier and not isinstance(value, str):
            issues.append(field + "：编号以数值存储，须核对前导零及精度")
        return str(value).strip()

    def amount(field, required=False):
        value = raw(field)
        if value is None:
            if required:
                issues.append(field + "：来源缺失")
            return None
        try:
            return money_value(value)
        except ValueError as exc:
            issues.append(field + "：" + str(exc))
            return None

    def date_field(field, required=False):
        value = raw(field)
        if value is None:
            if required:
                issues.append(field + "：来源缺失")
            return None
        try:
            return date_value(value)
        except ValueError as exc:
            issues.append(field + "：" + str(exc))
            return None

    def boolean_field(field, required=False):
        value = raw(field)
        if value is None:
            if required:
                issues.append(field + "：来源缺失")
            return None
        normalized = str(value).strip().lower()
        if normalized in {"是", "有", "需要", "true", "1", "yes", "y"}:
            return True
        if normalized in {"否", "无", "不需要", "false", "0", "no", "n"}:
            return False
        issues.append(field + "：只能明确填写是/否，不能猜测")
        return None

    definitions = GENERIC_FIELDS[kind]
    required = {
        "payroll": {"person_name", "actual_salary"},
        "social_security": {"person_id", "period_ref", "insurance_type", "employer_amount", "employee_amount"},
        "housing_fund": {"account", "amount", "entry_date"},
        "electronic_acceptance": {"acceptance_no", "amount", "transaction_date"},
        "individual_income_tax": {"person_name", "income"},
        "opening_balance": {"account_code", "account_name", "closing_debit", "closing_credit", "opening_debit", "opening_credit", "requires_auxiliary"},
        "contract": {"contract_no", "contract_date", "supplier", "amount"},
        "stock_in": {"stock_in_no", "stock_in_date", "supplier", "amount"},
    }[kind]
    money_fields = {
        "payroll": {"basic_salary", "actual_salary", "tax", "employer_social", "employee_social", "housing_fund", "employer_housing_fund", "employee_housing_fund"},
        "social_security": {"base", "employer_amount", "employee_amount", "total"},
        "housing_fund": {"amount", "subsidy"},
        "electronic_acceptance": {"amount"},
        "individual_income_tax": {"income", "expense", "tax_free_income", "pension", "medical", "unemployment", "housing_fund"},
        "opening_balance": {"closing_debit", "closing_credit", "opening_debit", "opening_credit"},
        "contract": {"amount"}, "stock_in": {"amount"},
    }[kind]
    date_fields = {
        "social_security": {"arrival_date"}, "housing_fund": {"entry_date"},
        "electronic_acceptance": {"issue_date", "maturity_date", "transaction_date"},
        "payroll": set(), "individual_income_tax": set(),
        "opening_balance": set(),
        "contract": {"contract_date"}, "stock_in": {"stock_in_date"},
    }[kind]
    identifier_fields = {"person_id", "account", "acceptance_no", "account_code", "contract_no", "stock_in_no"}
    normalized = {}
    row_period = None
    for field in definitions:
        if field in money_fields:
            normalized[field] = amount(field, required=field in required)
        elif field in date_fields:
            normalized[field] = date_field(field, required=field in required)
            if field in {"arrival_date", "entry_date", "transaction_date", "contract_date", "stock_in_date"} and normalized[field]:
                row_period = normalized[field][:7]
        elif field == "period_ref" or field in {"start_period", "end_period"}:
            value = raw(field)
            if value is None:
                normalized[field] = None
                if field in required:
                    issues.append(field + "：来源缺失")
            else:
                try:
                    normalized[field] = period_value(value)
                    if field == "period_ref" or row_period is None:
                        row_period = normalized[field]
                except ValueError as exc:
                    normalized[field] = None
                    issues.append(field + "：" + str(exc))
        elif field == "requires_auxiliary":
            normalized[field] = boolean_field(field, required=field in required)
        else:
            normalized[field] = text(field, required=field in required, identifier=field in identifier_fields)
    if row_period is None:
        # Opening-balance sources are uploaded as prior-period exceptions. If
        # the workbook has no period cell, the already-validated artifact
        # period is the only acceptable fallback; it is still checked by the
        # baseline candidate flow before confirmation.
        row_period = document_period or (period if kind == "opening_balance" else None)
    if payroll_context:
        fields["period"] = payroll_context["period_source"]
        row_period = fields["period"]["value"]
        if fields["period"]["status"] == "AMBIGUOUS":
            issues.append("period：工资期间标题冲突、非法或含公式，须核对具体来源")
        for index, reason in payroll_context["unresolved"].items():
            field = f"unmapped_{column_name(index+1)}"
            fields[field] = {**cell_source(sheet, row, headers, index, payroll_context["header_paths"]), "status": "AMBIGUOUS"}
            issues.append(f"{field}：{reason}（{fields[field]['region']}）")
    normalized["period"] = row_period
    if row_period is None:
        issues.append("period：原件没有可核验的业务期间")
    normalized["bank_account_ref"] = bank_account if kind == "electronic_acceptance" else None
    return {"record_type": GENERIC_RECORD_TYPES[kind], "source_anchor": {"row": number, "region": f"{sheet}!A{number}:{column_name(len(headers))}{number}"},
            "original_value": {"sheet": sheet, "row": number, "headers": headers, "values": values, "cell_types": row["types"], "formulas": row["formulas"],
                               **({"header_rows": payroll_context["header_rows"]} if payroll_context else {})},
            "normalized_value": normalized, "field_sources": fields, "extraction_issues": list(dict.fromkeys(issues)),
            "extraction_confidence": 0.0 if issues else 1.0,
            "period_check": "PASS" if row_period == period else "PERIOD_EXCEPTION"}


def extract_row(sheet, row, headers, mapping, kind, period, bank_account):
    if kind in GENERIC_RECORD_TYPES:
        return extract_generic_row(sheet, row, headers, mapping, kind, period, bank_account, None)
    values, number, issues, fields = row["values"], row["row"], [], {}
    def raw(field):
        for index in mapping.get(field, []):
            value = values[index] if index < len(values) else None
            fields[field] = cell_source(sheet, row, headers, index)
            if column_name(index+1)+str(number) in row["formulas"]:
                issues.append(field + "：公式值须与原件计算结果核实")
            if value not in (None, ""):
                return value
        return None
    def text(field, required=False, identifier=False):
        value = raw(field)
        if value is None:
            if required:
                issues.append(field + "：来源缺失")
            return None
        if identifier and not isinstance(value, str):
            issues.append(field + "：编号以数值存储，须核对前导零及精度")
        return str(value).strip()
    def amount(field, optional=False):
        value = raw(field)
        if optional and value is None:
            return None
        try:
            return money_value(value)
        except ValueError as exc:
            issues.append(field + "：" + str(exc))
            return None
    date_field = "transaction_date" if kind == "bank_statement" else "invoice_date"
    try:
        day = date_value(raw(date_field))
    except ValueError:
        day = None
        issues.append(date_field + "：日期缺失或不能确定")
    normalized = {date_field: day, "period": day[:7] if day else None}
    if kind != "bank_statement":
        record_type = "INVOICE" if kind == "purchase_invoices" else "SALES_INVOICE"
        normalized.update({field: text(field, required=field == "invoice_no", identifier=field.endswith("_id") or field == "invoice_no")
                           for field in ("invoice_no", "seller_name", "seller_tax_id", "buyer_name", "buyer_tax_id", "invoice_status")})
        normalized.update({field: amount(field) for field in ("net_amount", "tax")})
        total = amount("invoice_total", optional=True)
        if not mapping["invoice_total"] and all(normalized[f] is not None for f in ("net_amount", "tax")):
            total = format(Decimal(normalized["net_amount"]) + Decimal(normalized["tax"]), ".2f")
            normalized["invoice_total_derivation"] = "net_amount + tax"
            fields["invoice_total"] = {"row": number, "region": f"{sheet}!{column_name(mapping['net_amount'][0]+1)}{number}:{column_name(mapping['tax'][0]+1)}{number}"}
        normalized["invoice_total"] = total
        if total is None:
            issues.append("invoice_total：价税合计缺失或无法计算")
        if total is not None and all(normalized[f] is not None for f in ("net_amount", "tax")):
            if Decimal(total) != Decimal(normalized["net_amount"]) + Decimal(normalized["tax"]):
                issues.append("invoice_total：价税合计与金额、税额不一致")
            if Decimal(total) <= 0:
                issues.append("invoice_total：红字或零金额发票须人工核对")
        rate = raw("tax_rate")
        normalized["tax_rate"] = None
        rate_status = "MISSING"
        fields.setdefault("tax_rate", {"row": number, "region": "", "label": "税率", "source_label": "税率", "header": "税率",
                                       "original_value": None, "raw_value": None, "rawvalue": None, "header_path": []})
        try:
            if rate is None or (isinstance(rate, str) and not rate.strip()):
                pass  # An absent rate is not an extraction failure and is never inferred.
            elif isinstance(rate, bool):
                raise ValueError()
            else:
                rate_text = str(rate).strip()
                numeric_rate = Decimal(rate_text.removesuffix("%")) / (100 if rate_text.endswith("%") else 1)
                if not numeric_rate.is_finite() or not 0 <= numeric_rate <= 1:
                    raise ValueError()
                normalized["tax_rate"] = format(numeric_rate, "f")
                rate_status = "EXPLICIT"
        except (InvalidOperation, ValueError):
            rate_status = "INVALID"
            issues.append("tax_rate：原件税率非法，须核对；不能由税额比例推定")
        fields["tax_rate"]["status"] = rate_status
        if normalized["tax_rate"] is not None and all(normalized[f] is not None for f in ("net_amount", "tax")):
            expected_tax = (Decimal(normalized["net_amount"]) * Decimal(normalized["tax_rate"])).quantize(Decimal(".01"), rounding=ROUND_HALF_UP)
            if abs(Decimal(normalized["tax"]) - expected_tax) > Decimal(".01"):
                issues.append(f"tax：原件税额与明确税率不一致（按率校验为{expected_tax:.2f}，保留原税额）")
        if normalized["invoice_status"] not in {None, "正常", "有效"}:
            issues.append("invoice_status：发票状态须人工核对")
        normalized["counterparty"] = normalized["seller_name" if kind == "purchase_invoices" else "buyer_name"]
    else:
        normalized.update({field: text(field, identifier=field in {"transaction_id", "counterparty_account"})
                           for field in ("transaction_id", "counterparty", "counterparty_account", "summary", "note")})
        # A blank opposite-side amount is a structural zero in bank exports;
        # malformed non-blank values must remain unknown and never become zero.
        expense_raw, income_raw = raw("expense"), raw("income")
        def bank_amount(field, value):
            if value is None:
                return "0.00"
            try:
                return money_value(value)
            except ValueError as exc:
                issues.append(field + "：" + str(exc))
                return None
        expense, income = bank_amount("expense", expense_raw), bank_amount("income", income_raw)
        normalized.update({"expense": expense, "income": income, "balance": amount("balance", optional=True), "bank_account_ref": bank_account})
        if not bank_account:
            issues.append("bank_account_ref：本方账户待人工核对")
        record_type = "BANK_TRANSACTION"
        if expense is not None and income is not None and Decimal(expense) > 0 and Decimal(income) == 0:
            record_type = "PAYMENT"
            normalized["payment_total"] = expense
        elif expense is not None and income is not None and Decimal(income) > 0 and Decimal(expense) == 0:
            record_type = "RECEIPT"
            normalized["receipt_total"] = income
        else:
            issues.append("direction：收入支出方向不明确，须核对冲正或汇总记录")
    return {"record_type": record_type, "source_anchor": {"row": number, "region": f"{sheet}!A{number}:{column_name(len(headers))}{number}"},
            "original_value": {"sheet": sheet, "row": number, "headers": headers, "values": values, "cell_types": row["types"], "formulas": row["formulas"]},
            "normalized_value": normalized, "field_sources": fields, "extraction_issues": list(dict.fromkeys(issues)),
            "extraction_confidence": 0.0 if issues else 1.0,
            "period_check": "PASS" if day and day[:7] == period else "PERIOD_EXCEPTION"}


def bank_footer_header(values):
    labels = [str(v or '').strip() for v in values]
    return labels[:4] == ['总收入笔数', '总收入金额', '总支出笔数', '总支出金额'] and not any(labels[4:])


def bank_footer_values(values):
    if len(values) < 4 or any(v not in (None, '') for v in values[4:]):
        return False
    try:
        return (all(re.fullmatch(r'\d+(?:\.0+)?', str(values[i]).strip()) for i in (0, 2))
                and all(Decimal(money_value(values[i])) >= 0 for i in (1, 3)))
    except ValueError:
        return False


def extract_workbook(content, options: ParseOptions, period):
    from app.document_layouts import special_sheet, extract_boc_pdf
    if content.startswith(b'%PDF'):
        return extract_boc_pdf(content, options.document_kind, period)
    sheets = read_workbook(content)
    special = [special_sheet(s, options.document_kind, period, options.bank_account_ref) for s in sheets]
    if any(special):
        if len(sheets)!=1 or not all(special):
            raise ExtractionError('专用版式存在多表或未识别工作表，需先明确资料结构')
        return special[0]
    records, summaries = [], []
    document_period = detect_workbook_period(sheets)
    prefer_base = options.document_kind in {"purchase_invoices", "sales_invoices"} and any(s["name"] == "发票基础信息" for s in sheets)
    required_fields = {
        "payroll": ("person_name",), "social_security": ("person_id",), "housing_fund": ("account",),
        "electronic_acceptance": ("acceptance_no",), "individual_income_tax": ("person_name",),
        "opening_balance": ("account_code",), "contract": ("contract_no",), "stock_in": ("stock_in_no",),
    }
    for sheet in sheets:
        # Formula-only rows with no cached values still require field validation.
        nonempty = [r for r in sheet["rows"] if any(v not in (None, "") for v in r["values"]) or r.get("formulas")]
        if not nonempty:
            summaries.append({"sheet": sheet["name"], "status": "EMPTY", "rows": 0})
            continue
        if options.document_kind == "payroll":
            payroll = extract_payroll_sheet(sheet, period)
            records.extend(payroll["records"])
            summaries.extend(payroll["sheets"])
            continue
        if prefer_base and sheet["name"] == "信息汇总表":
            summaries.append({"sheet": sheet["name"], "status": "REFERENCE_ONLY", "reason": "已使用发票基础信息，明细汇总页不再次计数", "rows": len(nonempty)})
            continue
        header = next(((r, header_mapping(r["values"], options.document_kind)) for r in nonempty[:20] if header_mapping(r["values"], options.document_kind)), None)
        if header is None:
            summaries.append({"sheet": sheet["name"], "status": "UNRECOGNIZED", "reason": "未找到当前资料类型所需的明确列标题", "rows": len(nonempty)})
            continue
        row, mapping = header
        parsed = 0
        references = []
        footer_row = None
        for candidate in nonempty:
            if candidate["row"] <= row["row"]:
                continue
            first = str(next((v for v in candidate["values"] if v not in (None, "")), "")).strip()
            # This two-row footer uses completely different columns from transactions.
            # Keep its source coordinates and totals, never parse it as a bank entry.
            if options.document_kind == 'bank_statement' and bank_footer_header(candidate['values']):
                footer_row = candidate['row']
                references.append({'row': footer_row, 'reason': 'BANK_TOTAL_HEADER', 'values': candidate['values']})
                continue
            if footer_row is not None:
                if candidate['row'] == footer_row + 1 and bank_footer_values(candidate['values']):
                    references.append({'row': candidate['row'], 'reason': 'BANK_TOTAL_VALUES', 'values': candidate['values']})
                    footer_row = None
                    continue
                footer_row = None
            # Tax bureau exports label a separate totals row as 合计行. Only treat
            # it as reference when the mapped invoice identity/date are absent.
            if options.document_kind in {'purchase_invoices', 'sales_invoices'} and first == '合计行':
                identity_columns = mapping.get('invoice_no', []) + mapping.get('invoice_date', [])
                if all(i >= len(candidate['values']) or candidate['values'][i] in (None, '') for i in identity_columns):
                    references.append({'row': candidate['row'], 'reason': 'INVOICE_TOTAL', 'values': candidate['values']})
                    continue
            if first in {"合计", "总计", "小计", "本页合计", "期初余额", "期末余额"} or candidate["values"] == row["values"]:
                if options.document_kind in {'purchase_invoices', 'sales_invoices'} and first in {'合计', '总计', '小计', '本页合计'}:
                    identity_columns = mapping.get('invoice_no', []) + mapping.get('invoice_date', [])
                    if all(i >= len(candidate['values']) or candidate['values'][i] in (None, '', first) for i in identity_columns):
                        references.append({'row': candidate['row'], 'reason': 'SUMMARY_REFERENCE', 'values': candidate['values']})
                continue
            if options.document_kind in GENERIC_RECORD_TYPES:
                signal = next((candidate["values"][index] for field in required_fields[options.document_kind] for index in mapping.get(field, []) if index < len(candidate["values"]) and candidate["values"][index] not in (None, "")), None)
                if signal is None:
                    continue
                if options.document_kind == "individual_income_tax" and re.fullmatch(r"\d+(?:\.0)?", str(signal).strip()):
                    continue
                if options.document_kind == "payroll":
                    person_index = mapping["person_name"][0]
                    person = str(candidate["values"][person_index] or "").strip() if person_index < len(candidate["values"]) else ""
                    salary_index = mapping["actual_salary"][0]
                    salary = candidate["values"][salary_index] if salary_index < len(candidate["values"]) else None
                    # Payroll workbooks commonly place explanatory rows and a
                    # second insurance breakdown beneath each employee row.
                    # They are source material, not additional payroll facts.
                    try:
                        money_value(salary)
                    except ValueError:
                        continue
                    if not person or person.startswith("薪资") or person.endswith("有限公司"):
                        continue
                records.append(extract_generic_row(sheet["name"], candidate, row["values"], mapping, options.document_kind, period, options.bank_account_ref, document_period))
            else:
                records.append(extract_row(sheet["name"], candidate, row["values"], mapping, options.document_kind, period, options.bank_account_ref))
            parsed += 1
        summaries.append({"sheet": sheet["name"], "status": "EXTRACTED", "rows": parsed, "header_row": row["row"], "document_period": document_period, "reference_rows": references})
    extracted = {"records": records, "sheets": summaries,
                 "errors": [] if records else ["没有提取到可定位的记录，请核对资料类型、表头及文件内容"]}
    # Reconcile the full extraction, including out-of-period records. These are
    # parsing checks, not missing customer evidence or changes to fact values.
    from app.reconcile import reconcile
    extracted["checks"] = reconcile(extracted, options.document_kind)
    return extracted
