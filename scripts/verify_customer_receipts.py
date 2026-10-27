"""Opt-in real-model verification of invoice-first customer receipts and supplier payments in an isolated folder."""
from dataclasses import replace
from decimal import Decimal
from pathlib import Path
import argparse,json,sys
import openpyxl
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app.autonomous_runtime import run
from app.bookkeeping_profile import snapshot
from app.config import Settings
from app.review_loop import REVIEW_POLICY_VERSION,effective_entries
from app.stage23_pilot import sha256

def verify(result):
 d=result['delivery'];entries=[e for rows in effective_entries(d).values() for e in rows]
 def amount(code,side):return sum((Decimal(str(e[side])) for e in entries if e['code']==code),Decimal(0))
 checks={
  'mechanical':result['mechanical']['passed'],
  'independent_review':result['review']['verdict']=='pass' and not result['review'].get('defects'),
  'current_profile':d.get('bookkeeping_profile')=={k:snapshot()[k] for k in ('id','version','selection_version')},
  'uninvoiced_receipts_customer_a':amount('112201','credit')==1250 and amount('112201','debit')==0,
  'prior_invoice_customer_b':amount('112202','credit')==500 and amount('112202','debit')==0,
  'later_invoice_customer_d':amount('112204','debit')==400 and amount('112204','credit')==0,
  'loan_is_not_receivable':amount('224101','credit')==200,
  'no_duplicate_or_uninvoiced_revenue':amount('5001','credit')==400,
  'no_invented_vat':amount('222101','credit')==0,
  'bank_receipts_preserved':amount('1002','debit')==1950 and amount('1002','credit')==650,
  'uninvoiced_supplier_payment':amount('220201','debit')==300 and amount('220201','credit')==0,
  'prior_supplier_invoice_settlement':amount('220202','debit')==250 and amount('220202','credit')==0,
  'later_supplier_invoice':amount('220204','credit')==150 and amount('220204','debit')==0,
  'payment_loan_kept_separate':amount('122101','debit')==100,
  'only_invoiced_expense':amount('560202','debit')==150,
  'no_invented_input_vat':amount('222102','debit')==0,
  'no_unnecessary_question':not any(f.get('needs_user') for f in d['findings']),
  'all_candidates_reviewable':bool(d['voucher_assessments']) and all(a['status']=='READY_FOR_REVIEW' for a in d['voucher_assessments']),
  'draft_only':result['formal_balance_impact'] is False,
 }
 return {'passed':all(checks.values()),'checks':checks}

def main():
 p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);p.add_argument('--host-run-root',required=True);args=p.parse_args()
 out=args.output.resolve();out.mkdir(parents=True,exist_ok=False)
 book=openpyxl.Workbook();bank=book.active;bank.title='银行回单';bank.append(['日期','收入','支出','对方','摘要'])
 for r in [('2026-08-03',1000,0,'客户甲','业务货款'),('2026-08-05',250,0,'客户甲','追加业务货款'),('2026-08-08',500,0,'客户乙','7月已开票货款结清'),('2026-08-09',200,0,'企业丙','临时借款给本公司，非货款，9月归还')]:bank.append(r)
 for r in [('2026-08-15',0,300,'供应商甲','采购业务付款，暂未提供发票'),('2026-08-16',0,250,'供应商乙','结清7月已入账应付货款'),('2026-08-17',0,100,'企业戊','临时借给对方的款项，与采购无关，9月归还')]:bank.append(r)
 bills=book.create_sheet('销售发票');bills.append(['票号','日期','客户','金额','税额','价税合计','说明']);bills.append(['D0801','2026-08-12','客户丁',400,0,400,'免税普通票，本期首次开票；7月已收400挂应收贷方，本月无收款'])
 purchases=book.create_sheet('费用发票');purchases.append(['票号','日期','供应商','金额','税额','价税合计','说明']);purchases.append(['P0801','2026-08-18','供应商丁',150,0,150,'管理部门本月办公服务费，免税普通发票无抵扣进项，本期首次记费；7月已付款150挂应付借方，本月无付款'])
 chart=book.create_sheet('科目表');chart.append(['编码','名称'])
 for r in [('1002','银行存款'),('112201','应收账款—客户甲'),('112202','应收账款—客户乙'),('112204','应收账款—客户丁'),('224101','其他应付款—企业丙'),('5001','主营业务收入'),('222101','应交税费—销项税额'),('3103','本年利润'),('3001','实收资本')]:chart.append(r)
 for r in [('220201','应付账款—供应商甲'),('220202','应付账款—供应商乙'),('220204','应付账款—供应商丁'),('122101','其他应收款—企业戊'),('560202','管理费用—服务费'),('222102','应交税费—进项税额')]:chart.append(r)
 opening=book.create_sheet('已确认期初');opening.append(['科目编码','借方','贷方'])
 for r in [('1002',1000,0),('112201',0,0),('112202',500,0),('112204',0,400),('220204',150,0),('220202',0,250),('3001',0,1000)]:opening.append(r)
 facts=book.create_sheet('业务说明');facts.append(['说明'])
 for line in ['本期全部资料已经提供，无本期原账。客户甲正常业务来款，本期没有提供销售发票，期初应收为零；无需确认交货或验收。','客户乙7月已开票并确认收入500，期初应收借方500，本月收款结清，不重复确认销售。','企业丙200是借款，双方约定9月还本，本期无利息，无任何销售发票或销售关系。','客户丁7月收款400已贷记应收，本月仅补开票，8月无新收款；票面销售免税，本企业适用免税已确认，无销项税。','供应商甲期初为零，本月付款300暂未提供发票；供应商乙7月发票和费用已入账250，本月付款仅结清往来；供应商丁7月预先付款150已挂应付借方，本月取得办公服务发票首次记费，不能再造付款。企业戊100是明确的临时借出款，与供应商业务无关。除列明业务外，无工资、折旧或其他收入费用，销售服务无直接成本；结转本期有据损益即可。']:facts.append([line])
 source=out/'anonymous.xlsx';book.save(source)
 manifest={'period':'2026-08','sources':[{'id':'anonymous','filename':source.name,'path':'sources/'+source.name,'sha256':sha256(source)}]}
 context={'review_loop_version':REVIEW_POLICY_VERSION,'statements':[],'materials_complete':True,'bookkeeping_profile':snapshot(),'goal':'根据所有原件及期初按所选记账方式生成本期待复核草稿，不改前期账。'}
 settings=replace(Settings.from_env(),autonomous_backend='claude_code',autonomous_execution_mode='docker',database_path=out/'test.db',autonomous_host_run_root=args.host_run_root,autonomous_max_rounds=3)
 result=run(out/'autonomous-runs/run',manifest,{'anonymous':source},context,settings,lambda:None,lambda phase,round_no,message:print(phase,round_no,message,flush=True))
 (out/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));report=verify(result);(out/'checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False),flush=True)
 return 0 if report['passed'] else 1
if __name__=='__main__':raise SystemExit(main())
